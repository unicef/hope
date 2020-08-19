import json
from collections import defaultdict
from io import BytesIO
from typing import Union

import openpyxl
from django.contrib.gis.geos import Point
from django.core.files import File
from django.core.files.storage import default_storage
from django.core.management import call_command
from django.db import transaction
from django.utils import timezone
from django_countries.fields import Country
from openpyxl_image_loader import SheetImageLoader

from core.core_fields_attributes import (
    TYPE_INTEGER,
    TYPE_SELECT_ONE,
    COLLECTORS_FIELDS,
)
from core.kobo.api import KoboAPI
from core.kobo.common import get_field_name, KOBO_FORM_INDIVIDUALS_COLUMN_NAME
from core.models import BusinessArea
from core.utils import (
    get_combined_attributes,
    serialize_flex_attributes,
    rename_dict_keys,
)
from household.const import COUNTRIES_NAME_ALPHA2
from household.models import (
    IDENTIFICATION_TYPE_DICT,
    ROLE_PRIMARY,
    ROLE_ALTERNATE,
    YES,
    HEAD,
    NON_BENEFICIARY,
)
from registration_data.models import RegistrationDataImport
from registration_datahub.models import (
    ImportData,
    ImportedAgency,
    ImportedIndividualIdentity,
    ImportedIndividualRoleInHousehold,
)
from registration_datahub.models import (
    ImportedDocument,
    ImportedDocumentType,
)
from registration_datahub.models import (
    ImportedHousehold,
    ImportedIndividual,
)
from registration_datahub.models import RegistrationDataImportDatahub
from registration_datahub.tasks.batch_deduplicate import DeduplicateTask
from registration_datahub.tasks.utils import collectors_str_ids_to_list


class RdiBaseCreateTask:
    COMBINED_FIELDS = get_combined_attributes()
    FLEX_FIELDS = serialize_flex_attributes()

    def _cast_value(self, value, header):
        if value in (None, ""):
            return value

        value_type = self.COMBINED_FIELDS[header]["type"]

        if value_type == TYPE_INTEGER:
            return int(value)

        if value_type == TYPE_SELECT_ONE:
            custom_cast_method = self.COMBINED_FIELDS[header].get(
                "custom_cast_value"
            )

            if custom_cast_method is not None:
                return custom_cast_method(input_value=value)

            choices = [
                x.get("value") for x in self.COMBINED_FIELDS[header]["choices"]
            ]
            if isinstance(value, str):
                upper_value = value.upper()
                if upper_value in choices:
                    return upper_value

            if value not in choices:
                try:
                    return int(value)
                except ValueError:
                    return str(value)

        return value


class RdiXlsxCreateTask(RdiBaseCreateTask):
    """
    Works on valid XLSX files, parsing them and creating households/individuals
    in the Registration Datahub. Once finished it will update the status of
    that registration data import instance.
    """

    image_loader = None
    business_area = None
    households = None
    individuals = None
    documents = None
    identities = None
    collectors = None

    def _handle_document_fields(
        self, value, header, row_num, individual, *args, **kwargs
    ):
        if value is None:
            return

        if header.startswith("other_id"):
            document_data = self.documents.get(f"individual_{row_num}_other")
        else:
            document_data = self.documents.get(f"individual_{row_num}_{header}")

        if header == "other_id_type_i_c":
            if document_data:
                document_data["name"] = value
            else:
                self.documents[f"individual_{row_num}_other"] = {
                    "individual": individual,
                    "name": value,
                    "type": "OTHER",
                }
        elif header == "other_id_no_i_c":
            if document_data:
                document_data["value"] = value
            else:
                self.documents[f"individual_{row_num}_other"] = {
                    "individual": individual,
                    "value": value,
                    "type": "OTHER",
                }
        else:
            document_name = header.replace("_no", "").replace("_i_c", "")
            doc_type = document_name.upper().strip()

            if document_data:
                document_data["value"] = value
            else:
                self.documents[f"individual_{row_num}_{header}"] = {
                    "individual": individual,
                    "name": IDENTIFICATION_TYPE_DICT.get(doc_type),
                    "type": doc_type,
                    "value": value,
                }

    def _handle_document_photo_fields(
        self, cell, row_num, individual, header, *args, **kwargs
    ):
        if not self.image_loader.image_in(cell.coordinate):
            return

        if header.startswith("other_id"):
            document_data = self.documents.get(f"individual_{row_num}_other")
        else:
            document_data = self.documents.get(f"individual_{row_num}_{header}")

        file = self._handle_image_field(cell)

        if document_data:
            document_data["photo"] = file
        else:
            suffix = (
                "other"
                if header.startswith("other_id")
                else header.replace("photo", "no")
            )
            self.documents[f"individual_{row_num}_{suffix}"] = {
                "individual": individual,
                "photo": file.name,
            }

    def _handle_image_field(self, cell, is_flex_field=False, *args, **kwargs):
        if self.image_loader.image_in(cell.coordinate):
            image = self.image_loader.get(cell.coordinate)
            file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
            file_io = BytesIO()

            image.save(file_io, image.format)

            file = File(file_io, name=file_name)

            if is_flex_field:
                return default_storage.save(file_name, file)

            return file
        return ""

    def _handle_geopoint_field(self, value, *args, **kwargs):
        if not value:
            return ""

        values_as_list = value.split(",")
        longitude = values_as_list[0].strip()
        latitude = values_as_list[1].strip()

        return Point(x=float(longitude), y=float(latitude), srid=4326)

    def _handle_identity_fields(
        self, value, header, row_num, individual, *args, **kwargs
    ):
        if value is None:
            return

        agency = ImportedAgency.objects.get(
            type="WFP" if header == "scope_id_no_i_c" else "UNHCR"
        )

        identities_data = self.identities.get(f"individual_{row_num}")

        if identities_data:
            identities_data["number"] = value
            identities_data["agency"] = agency

        self.identities[f"individual_{row_num}"] = {
            "individual": individual,
            "number": value,
            "agency": agency,
        }

    def _handle_identity_photo(
        self, cell, row_num, individual, *args, **kwargs
    ):
        if not self.image_loader.image_in(cell.coordinate):
            return

        identity_data = self.identities.get(f"individual_{row_num}")

        image = self.image_loader.get(cell.coordinate)
        file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
        file_io = BytesIO()

        image.save(file_io, image.format)

        file = File(file_io, name=file_name)

        if identity_data:
            identity_data["photo"] = file
        else:
            self.identities[f"individual_{row_num}"] = {
                "individual": individual,
                "photo": file_name,
            }

    def _handle_collectors(self, value, header, individual, *args, **kwargs):
        list_of_ids = collectors_str_ids_to_list(value)
        if list_of_ids is None:
            return

        for hh_id in list_of_ids:
            role = (
                ROLE_PRIMARY
                if header == "primary_collector_id"
                else ROLE_ALTERNATE
            )
            self.collectors[hh_id].append(
                ImportedIndividualRoleInHousehold(
                    individual=individual, role=role
                )
            )

    def _create_documents(self):
        docs_to_create = []
        for document_data in self.documents.values():
            doc_type, is_created = ImportedDocumentType.objects.get_or_create(
                country=Country(
                    COUNTRIES_NAME_ALPHA2.get(self.business_area.name.title())
                ),
                label=document_data["name"],
                type=document_data["type"],
            )
            photo = document_data.get("photo")
            individual = document_data.get("individual")
            obj = ImportedDocument(
                document_number=document_data.get("value"),
                photo=photo,
                individual=individual,
                type=doc_type,
            )

            docs_to_create.append(obj)

        ImportedDocument.objects.bulk_create(docs_to_create)

    def _create_identities(self):
        idents_to_create = [
            ImportedIndividualIdentity(
                agency=ident_data["agency"],
                individual=ident_data["individual"],
                document_number=ident_data["number"],
            )
            for ident_data in self.identities.values()
        ]

        ImportedIndividualIdentity.objects.bulk_create(idents_to_create)

    def _create_collectors(self):
        collectors_to_create = []
        for hh_id, collectors_list in self.collectors.items():
            for collector in collectors_list:
                collector.household_id = self.households.get(hh_id).pk
                collectors_to_create.append(collector)
        ImportedIndividualRoleInHousehold.objects.bulk_create(
            collectors_to_create
        )

    def _create_objects(self, sheet, registration_data_import):
        complex_fields = {
            "individuals": {
                "birth_certificate_no_i_c": self._handle_document_fields,
                "birth_certificate_photo_i_c": self._handle_document_photo_fields,
                "drivers_license_no_i_c": self._handle_document_fields,
                "drivers_license_photo_i_c": self._handle_document_photo_fields,
                "electoral_card_no_i_c": self._handle_document_fields,
                "electoral_card_photo_i_c": self._handle_document_photo_fields,
                "unhcr_id_no_i_c": self._handle_identity_fields,
                "national_id_no_ic": self._handle_document_fields,
                "national_id_photo_ic": self._handle_document_photo_fields,
                "national_passport_i_c": self._handle_document_fields,
                "national_passport_photo_i_c": self._handle_document_photo_fields,
                "scope_id_no_i_c": self._handle_identity_fields,
                "other_id_type_i_c": self._handle_document_fields,
                "other_id_no_i_c": self._handle_document_fields,
                "other_id_photo_i_c": self._handle_document_photo_fields,
                "photo_i_c": self._handle_image_field,
                "primary_collector_id": self._handle_collectors,
                "alternate_collector_id": self._handle_collectors,
            },
            "households": {
                "consent_h_c": self._handle_image_field,
                "hh_geopoint_h_c": self._handle_geopoint_field,
            },
        }

        complex_types = {
            "GEOPOINT": self._handle_geopoint_field,
            "IMAGE": self._handle_image_field,
        }

        sheet_title = sheet.title.lower()

        first_row = sheet[1]
        households_to_update = []
        for row in sheet.iter_rows(min_row=3):
            if not any([cell.value for cell in row]):
                continue

            if sheet_title == "households":
                obj_to_create = ImportedHousehold(
                    registration_data_import=registration_data_import,
                )
            else:
                obj_to_create = ImportedIndividual(
                    registration_data_import=registration_data_import,
                )

            household_id = None
            for cell, header_cell in zip(row, first_row):
                header = header_cell.value
                combined_fields = {**self.COMBINED_FIELDS, **COLLECTORS_FIELDS}
                current_field = combined_fields.get(header)

                if not current_field:
                    continue

                is_not_image = current_field["type"] != "IMAGE"

                is_not_required_and_empty = (
                    not current_field.get("required")
                    and cell.value is None
                    and is_not_image
                )

                if is_not_required_and_empty:
                    continue

                if header == "household_id":
                    temp_value = cell.value
                    if (
                        isinstance(temp_value, float)
                        and temp_value.is_integer()
                    ):
                        temp_value = int(temp_value)
                    household_id = str(temp_value)
                    if sheet_title == "individuals":
                        obj_to_create.household = self.households.get(
                            household_id
                        )

                if header in complex_fields[sheet_title]:
                    fn = complex_fields[sheet_title].get(header)
                    value = fn(
                        value=cell.value,
                        cell=cell,
                        header=header,
                        row_num=cell.row,
                        individual=obj_to_create
                        if sheet_title == "individuals"
                        else None,
                    )
                    if value is not None:
                        setattr(
                            obj_to_create,
                            combined_fields[header]["name"],
                            value,
                        )
                elif (
                    hasattr(obj_to_create, combined_fields[header]["name"],)
                    and header != "household_id"
                ):

                    value = self._cast_value(cell.value, header)
                    if value in (None, ""):
                        continue

                    if header == "relationship_i_c" and value == HEAD:
                        household = self.households.get(household_id)
                        if household is not None:
                            household.head_of_household = obj_to_create
                            households_to_update.append(household)

                    setattr(
                        obj_to_create, combined_fields[header]["name"], value,
                    )
                elif header in self.FLEX_FIELDS[sheet_title]:
                    value = self._cast_value(cell.value, header)
                    type_name = self.FLEX_FIELDS[sheet_title][header]["type"]
                    if type_name in complex_types:
                        fn = complex_types[type_name]
                        value = fn(
                            value=cell.value,
                            cell=cell,
                            header=header,
                            is_flex_field=True,
                        )
                    obj_to_create.flex_fields[header] = value

            obj_to_create.last_registration_date = (
                obj_to_create.first_registration_date
            )

            if sheet_title == "households":
                self.households[household_id] = obj_to_create
            else:
                if household_id is None:
                    obj_to_create.relationship = NON_BENEFICIARY
                self.individuals.append(obj_to_create)

        if sheet_title == "households":
            ImportedHousehold.objects.bulk_create(self.households.values())
        else:
            ImportedIndividual.objects.bulk_create(self.individuals)
            ImportedHousehold.objects.bulk_update(
                households_to_update, ["head_of_household"], 1000,
            )
            self._create_documents()
            self._create_identities()
            self._create_collectors()

    @transaction.atomic(using="default")
    @transaction.atomic(using="registration_datahub")
    def execute(
        self, registration_data_import_id, import_data_id, business_area_id
    ):
        self.households = {}
        self.documents = {}
        self.identities = {}
        self.individuals = []
        self.collectors = defaultdict(list)

        registration_data_import = RegistrationDataImportDatahub.objects.select_for_update().get(
            id=registration_data_import_id,
        )
        registration_data_import.import_done = (
            RegistrationDataImportDatahub.STARTED
        )
        registration_data_import.save()

        import_data = ImportData.objects.get(id=import_data_id)

        self.business_area = BusinessArea.objects.get(id=business_area_id)

        wb = openpyxl.load_workbook(import_data.file, data_only=True)

        # households objects have to be create first
        worksheets = (wb["Households"], wb["Individuals"])
        for sheet in worksheets:
            self.image_loader = SheetImageLoader(sheet)
            self._create_objects(sheet, registration_data_import)

        registration_data_import.import_done = (
            RegistrationDataImportDatahub.DONE
        )
        registration_data_import.save()

        RegistrationDataImport.objects.filter(
            id=registration_data_import.hct_id
        ).update(status=RegistrationDataImport.IN_REVIEW)

        call_command(
            "search_index",
            "--populate",
            "--models",
            "registration_datahub.ImportedIndividual",
        )

        DeduplicateTask.deduplicate_imported_individuals(
            registration_data_import_datahub=registration_data_import
        )


class RdiKoboCreateTask(RdiBaseCreateTask):
    """
    Imports project data from Kobo via a REST API, parsing them and creating
    households/individuals in the Registration Datahub. Once finished it will
    update the status of that registration data import instance.
    """

    DOCS_AND_IDENTITIES_FIELDS = {
        "birth_certificate_no_i_c",
        "birth_certificate_photo_i_c",
        "drivers_license_no_i_c",
        "drivers_license_photo_i_c",
        "electoral_card_no_i_c",
        "electoral_card_photo_i_c",
        "unhcr_id_no_i_c",
        "unhcr_id_photo_i_c",
        "national_id_no_ic",
        "national_id_photo_ic",
        "national_passport_i_c",
        "national_passport_photo_i_c",
        "scope_id_no_i_c",
        "scope_id_photo_i_c",
        "other_id_type_i_c",
        "other_id_no_i_c",
        "other_id_photo_i_c",
    }

    reduced_submissions = None
    business_area = None
    attachments = None

    def _handle_image_field(self, value):
        download_url = ""
        for attachment in self.attachments:
            current_download_url = attachment.get("download_url", "")
            if current_download_url.endswith(value):
                download_url = current_download_url

        if not download_url:
            return download_url

        api = KoboAPI(self.business_area.slug)
        image_bytes = api.get_attached_file(download_url)
        file = File(image_bytes, name=value)

        return file

    def _handle_geopoint_field(self, value):
        geopoint = value.split(" ")
        x = float(geopoint[0])
        y = float(geopoint[1])
        return Point(x=x, y=y, srid=4326)

    def _cast_and_assign(
        self, value: Union[str, list], field: str, obj: object
    ):
        complex_fields = {
            "IMAGE": self._handle_image_field,
            "GEOPOINT": self._handle_geopoint_field,
        }
        excluded = ("age",)

        field_data_dict = self.COMBINED_FIELDS.get(field)

        if field_data_dict is None or field in excluded:
            return

        if field_data_dict["type"] in complex_fields:
            cast_fn = complex_fields.get(field_data_dict["type"])
            correct_value = cast_fn(value)
        else:
            correct_value = self._cast_value(value, field)

        setattr(obj, field_data_dict["name"], correct_value)

    def _handle_documents_and_identities(
        self, documents_and_identities, individual_dict
    ):
        identity_fields = {
            "scope_id",
            "unhcr_id",
        }

        documents = []
        identities = []
        for documents_dict in documents_and_identities:
            for document_name, data in documents_dict.items():
                is_identity = document_name in identity_fields

                if is_identity:
                    agency = ImportedAgency.objects.get(
                        type="WFP" if document_name == "scope_id" else "UNHCR"
                    )
                    identities.append(
                        ImportedIndividualIdentity(
                            agency=agency,
                            individual=individual_dict.get(
                                data["individual"].get_hash_key
                            ),
                            document_number=data["number"],
                        )
                    )
                else:
                    type_name = document_name.upper()
                    label = IDENTIFICATION_TYPE_DICT.get(type_name)
                    if label is None:
                        label = data["name"]
                    (
                        document_type,
                        is_created,
                    ) = ImportedDocumentType.objects.get_or_create(
                        country=Country(
                            COUNTRIES_NAME_ALPHA2.get(
                                self.business_area.name.title()
                            )
                        ),
                        label=label,
                        type=type_name,
                    )
                    file = self._handle_image_field(data.get("photo", ""))
                    documents.append(
                        ImportedDocument(
                            document_number=data["number"],
                            photo=file,
                            individual=data["individual"],
                            type=document_type,
                        )
                    )
        ImportedDocument.objects.bulk_create(documents)
        ImportedIndividualIdentity.objects.bulk_create(identities)

    @staticmethod
    def _handle_collectors(collectors_dict, individuals_dict):
        collectors_to_bulk_create = []
        for hash_key, collectors_list in collectors_dict.items():
            for collector in collectors_list:
                collector.individual = individuals_dict.get(hash_key)
                collectors_to_bulk_create.append(collector)
        ImportedIndividualRoleInHousehold.objects.bulk_create(
            collectors_to_bulk_create
        )

    @transaction.atomic(using="default")
    @transaction.atomic(using="registration_datahub")
    def execute(
        self, registration_data_import_id, import_data_id, business_area_id
    ):
        registration_data_import = RegistrationDataImportDatahub.objects.select_for_update().get(
            id=registration_data_import_id,
        )
        registration_data_import.import_done = (
            RegistrationDataImportDatahub.STARTED
        )
        registration_data_import.save()

        import_data = ImportData.objects.get(id=import_data_id)

        self.business_area = BusinessArea.objects.get(id=business_area_id)

        submissions_json = import_data.file.read()
        submissions = json.loads(submissions_json)
        self.reduced_submissions = rename_dict_keys(submissions, get_field_name)

        head_of_households_mapping = {}
        households_to_create = []
        individuals_to_create = {}
        documents_and_identities_to_create = []
        collectors_to_create = defaultdict(list)
        for household in self.reduced_submissions:
            collectors_count = 0
            household_obj = ImportedHousehold()
            self.attachments = household.get("_attachments", [])
            for hh_field, hh_value in household.items():
                self._cast_and_assign(hh_value, hh_field, household_obj)
                if hh_field == KOBO_FORM_INDIVIDUALS_COLUMN_NAME:
                    for individual in hh_value:
                        current_individual_docs_and_identities = defaultdict(
                            dict
                        )
                        individual_obj = ImportedIndividual()
                        only_collector_flag = False
                        role = None
                        for i_field, i_value in individual.items():
                            if i_field in self.DOCS_AND_IDENTITIES_FIELDS:
                                key = i_field.replace("_photo_i_c", "").replace(
                                    "_no_i_c", ""
                                )
                                if i_field.endswith("_type_i_c"):
                                    value_key = "name"
                                elif i_field.endswith("_photo_i_c"):
                                    value_key = "photo"
                                else:
                                    value_key = "number"
                                current_individual_docs_and_identities[key][
                                    value_key
                                ] = i_value
                                current_individual_docs_and_identities[key][
                                    "individual"
                                ] = individual_obj
                            elif i_field == "is_only_collector":
                                if i_value == YES:
                                    only_collector_flag = True
                                    collectors_count += 1
                            elif i_field == "role_i_c":
                                role = i_value.upper()
                            else:
                                self._cast_and_assign(
                                    i_value, i_field, individual_obj
                                )
                        if (
                            individual_obj.relationship == HEAD
                            and only_collector_flag is False
                        ):
                            head_of_households_mapping[
                                household_obj
                            ] = individual_obj

                        if only_collector_flag is True:
                            individual_obj.relationship = NON_BENEFICIARY

                        individual_obj.last_registration_date = (
                            individual_obj.first_registration_date
                        )
                        individual_obj.registration_data_import = (
                            registration_data_import
                        )

                        duplicated_object = individuals_to_create.get(
                            individual_obj.get_hash_key
                        )
                        has_documents = (
                            len(current_individual_docs_and_identities) > 0
                        )
                        if (
                            duplicated_object is None
                            or has_documents
                            or only_collector_flag is False
                        ):
                            individuals_to_create[
                                individual_obj.get_hash_key
                            ] = individual_obj

                        if only_collector_flag is True:
                            individual_obj.household = None
                        else:
                            individual_obj.household = household_obj

                        if role in (ROLE_PRIMARY, ROLE_ALTERNATE):
                            role_obj = ImportedIndividualRoleInHousehold(
                                individual=duplicated_object or individual_obj,
                                household_id=household_obj.pk,
                                role=role,
                            )
                            collectors_to_create[
                                individual_obj.get_hash_key
                            ].append(role_obj)

                        documents_and_identities_to_create.append(
                            current_individual_docs_and_identities
                        )

            household_obj.size = household_obj.size - collectors_count
            household_obj.last_registration_date = (
                household_obj.first_registration_date
            )
            household_obj.registration_data_import = registration_data_import
            households_to_create.append(household_obj)

        ImportedHousehold.objects.bulk_create(households_to_create)
        ImportedIndividual.objects.bulk_create(individuals_to_create.values())
        self._handle_collectors(collectors_to_create, individuals_to_create)
        self._handle_documents_and_identities(
            documents_and_identities_to_create, individuals_to_create,
        )

        households_to_update = []
        for household, individual in head_of_households_mapping.items():
            household.head_of_household = individual
            households_to_update.append(household)
        ImportedHousehold.objects.bulk_update(
            households_to_update, ["head_of_household"], 1000,
        )

        registration_data_import.import_done = (
            RegistrationDataImportDatahub.DONE
        )
        registration_data_import.save()

        RegistrationDataImport.objects.filter(
            id=registration_data_import.hct_id
        ).update(status=RegistrationDataImport.IN_REVIEW)

        call_command(
            "search_index",
            "--populate",
            "--models",
            "registration_datahub.ImportedIndividual",
        )

        DeduplicateTask.deduplicate_imported_individuals(
            registration_data_import_datahub=registration_data_import
        )
