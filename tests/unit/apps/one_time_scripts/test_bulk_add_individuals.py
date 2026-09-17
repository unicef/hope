from datetime import date
from typing import Any, Callable

import openpyxl
import pytest

from extras.test_utils.factories import (
    BusinessAreaFactory,
    FlexibleAttributeFactory,
    HouseholdFactory,
    IndividualFactory,
    ProgramFactory,
    UserFactory,
)
from hope.apps.grievance.models import GrievanceTicket
from hope.models import BusinessArea, FlexibleAttribute, Household, Individual, Program, User
from hope.one_time_scripts.bulk_add_individuals import bulk_add_individuals

pytestmark = [
    pytest.mark.usefixtures("mock_elasticsearch"),
    pytest.mark.django_db,
]

DEFAULT_HEADER = ("programme_name", "unicef_id", "full_name_i_c", "gender_i_c", "birth_date_i_c", "relationship_i_c")


@pytest.fixture
def business_area() -> BusinessArea:
    return BusinessAreaFactory(slug="myanmar", name="Myanmar")


@pytest.fixture
def program(business_area: BusinessArea) -> Program:
    return ProgramFactory(business_area=business_area, name="MCCT")


@pytest.fixture
def household(program: Program) -> Household:
    return HouseholdFactory(program=program, business_area=program.business_area, size=1, create_role=False)


@pytest.fixture
def other_household(program: Program) -> Household:
    return HouseholdFactory(program=program, business_area=program.business_area, size=1, create_role=False)


@pytest.fixture
def removed_household(program: Program) -> Household:
    return HouseholdFactory(
        program=program, business_area=program.business_area, size=1, create_role=False, is_removed=True
    )


@pytest.fixture
def pending_household(program: Program) -> Household:
    return HouseholdFactory(
        program=program, business_area=program.business_area, size=1, create_role=False, rdi_merge_status="PENDING"
    )


@pytest.fixture
def mother(household: Household) -> Individual:
    return IndividualFactory(household=household)


@pytest.fixture
def user() -> User:
    return UserFactory(email="loader@unicef.org")


@pytest.fixture
def mother_name_flex_field() -> FlexibleAttribute:
    return FlexibleAttributeFactory(
        name="ind_mother_full_name_i_f",
        type=FlexibleAttribute.STRING,
        associated_with=FlexibleAttribute.ASSOCIATED_WITH_INDIVIDUAL,
    )


@pytest.fixture
def birth_weight_flex_field() -> FlexibleAttribute:
    return FlexibleAttributeFactory(
        name="birth_weight_grams_i_f",
        type=FlexibleAttribute.INTEGER,
        associated_with=FlexibleAttribute.ASSOCIATED_WITH_INDIVIDUAL,
    )


@pytest.fixture
def vaccination_date_flex_field() -> FlexibleAttribute:
    return FlexibleAttributeFactory(
        name="vaccination_date_i_f",
        type=FlexibleAttribute.DATE,
        associated_with=FlexibleAttribute.ASSOCIATED_WITH_INDIVIDUAL,
    )


@pytest.fixture
def write_workbook(tmp_path: Any) -> Callable[..., str]:
    def write(rows: list[tuple], header: tuple = DEFAULT_HEADER) -> str:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(list(header))
        for row in rows:
            worksheet.append(list(row))
        path = tmp_path / "newborns.xlsx"
        workbook.save(path)
        return str(path)

    return write


def test_row_creates_individual_in_the_household(
    household: Household, program: Program, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert result["created"] == [2]
    assert individual.sex == "MALE"
    assert individual.birth_date == date(2026, 5, 12)
    assert individual.relationship == "SON_DAUGHTER"
    assert individual.program_id == program.pk
    assert individual.rdi_merge_status == "MERGED"
    assert individual.registration_data_import_id == household.registration_data_import_id


def test_closed_ticket_is_recorded_for_the_created_individual(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    ticket = GrievanceTicket.objects.get(household_unicef_id=household.unicef_id)
    assert ticket.status == GrievanceTicket.STATUS_CLOSED
    assert ticket.category == GrievanceTicket.CATEGORY_DATA_CHANGE
    assert ticket.issue_type == GrievanceTicket.ISSUE_TYPE_DATA_CHANGE_ADD_INDIVIDUAL
    assert ticket.add_individual_ticket_details.approve_status is True
    assert ticket.created_by == user


def test_household_size_grows_by_the_number_of_created_individuals(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby One", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Two", "FEMALE", date(2026, 6, 1), "SON_DAUGHTER"),
        ]
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    household.refresh_from_db()
    assert household.size == 3


def test_individual_count_is_reported_before_and_after(
    household: Household, user: User, write_workbook: Callable[..., str], capsys: pytest.CaptureFixture[str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby One", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Two", "FEMALE", date(2026, 6, 1), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["individuals_before"] == 1
    assert result["individuals_after"] == 3
    assert "individuals: 1 before, 3 after, +2 (OK, 2 rows created)" in capsys.readouterr().out


def test_validate_only_run_writes_nothing(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["created"] == [2]
    assert not Individual.objects.filter(full_name="Baby Aye").exists()
    assert not GrievanceTicket.objects.exists()


def test_flex_field_column_is_stored_on_the_individual(
    household: Household, user: User, mother_name_flex_field: FlexibleAttribute, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", "Daw Mya")],
        header=(*DEFAULT_HEADER, "ind_mother_full_name_i_f"),
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.flex_fields["ind_mother_full_name_i_f"] == "Daw Mya"


@pytest.mark.parametrize(
    "cell_value",
    [pytest.param("3200", id="text_cell"), pytest.param(3200, id="number_cell")],
)
def test_integer_flex_field_is_stored_as_a_number(
    household: Household,
    user: User,
    birth_weight_flex_field: FlexibleAttribute,
    write_workbook: Callable[..., str],
    cell_value: Any,
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", cell_value)],
        header=(*DEFAULT_HEADER, "birth_weight_grams_i_f"),
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.flex_fields["birth_weight_grams_i_f"] == 3200


def test_date_flex_field_is_stored_as_iso_text(
    household: Household,
    user: User,
    vaccination_date_flex_field: FlexibleAttribute,
    write_workbook: Callable[..., str],
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", date(2026, 6, 20))],
        header=(*DEFAULT_HEADER, "vaccination_date_i_f"),
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.flex_fields["vaccination_date_i_f"] == "2026-06-20"


def test_unregistered_flex_field_fails_the_row(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", "something")],
        header=(*DEFAULT_HEADER, "ind_unknown_i_f"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert "ind_unknown_i_f is not a correct" in result["failed"][0][1]
    assert not Individual.objects.filter(full_name="Baby Aye").exists()


@pytest.mark.parametrize(
    ("full_name", "sex", "birth_date", "relationship", "expected_error"),
    [
        pytest.param("", "MALE", date(2026, 5, 12), "SON_DAUGHTER", "missing required full_name", id="no_name"),
        pytest.param("Baby Aye", "", date(2026, 5, 12), "SON_DAUGHTER", "missing required sex", id="no_sex"),
        pytest.param("Baby Aye", "MALE", "", "SON_DAUGHTER", "missing required birth_date", id="no_birth_date"),
        pytest.param("Baby Aye", "MALE", date(2026, 5, 12), "", "missing required relationship", id="no_relationship"),
        pytest.param("Baby Aye", "Girl", date(2026, 5, 12), "SON_DAUGHTER", "invalid 'Girl'", id="bad_sex"),
        pytest.param("Baby Aye", "MALE", date(2026, 5, 12), "CHILD", "invalid 'CHILD'", id="bad_relationship"),
        pytest.param("Baby Aye", "MALE", "not a date", "SON_DAUGHTER", "does not match format", id="bad_birth_date"),
    ],
)
def test_invalid_row_fails_and_creates_nothing(
    household: Household,
    user: User,
    write_workbook: Callable[..., str],
    full_name: str,
    sex: str,
    birth_date: Any,
    relationship: str,
    expected_error: str,
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, full_name, sex, birth_date, relationship)])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert expected_error in result["failed"][0][1]
    assert Individual.objects.filter(household=household).count() == 1


def test_unexpected_error_is_reported_with_its_type(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", date(2026, 1, 1))],
        header=(*DEFAULT_HEADER, "age_at_registration"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["created"] == []
    assert result["failed"][0][0] == 2
    assert result["failed"][0][1].startswith("TypeError: int() argument")


def test_head_relationship_is_reported_as_a_warning(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "New Head", "MALE", date(1990, 5, 12), "HEAD")])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["created"] == [2]
    assert result["warnings"] == [
        (2, "relationship HEAD makes this individual the head of household and sets every other member to UNKNOWN")
    ]
    assert not Individual.objects.filter(full_name="New Head").exists()


def test_unknown_household_fails_only_that_row(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", "HH-00-0000.0000", "Missing Home", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Aye", "FEMALE", date(2026, 6, 1), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["created"] == [3]
    assert result["failed"] == [(2, "household HH-00-0000.0000 not found in programme MCCT")]


def test_apply_run_writes_nothing_when_a_reference_is_broken(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", "HH-00-0000.0000", "Missing Home", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Aye", "FEMALE", date(2026, 6, 1), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert result["failed"] == [(2, "household HH-00-0000.0000 not found in programme MCCT")]
    assert not Individual.objects.filter(full_name="Baby Aye").exists()
    assert not GrievanceTicket.objects.exists()


def test_removed_household_fails_the_row(
    removed_household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", removed_household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["failed"] == [(2, f"household {removed_household.unicef_id} removed in programme MCCT")]


def test_pending_household_fails_the_row(
    pending_household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", pending_household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["failed"] == [(2, f"household {pending_household.unicef_id} not merged in programme MCCT")]


def test_unknown_programme_fails_the_row(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook([("Nutrition", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert result["failed"] == [(2, "programme 'Nutrition' not found")]


def test_member_of_the_household_passes_the_row(
    household: Household, mother: Individual, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", mother.unicef_id)],
        header=(*DEFAULT_HEADER, "individual_unicef_id"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [2]
    assert Individual.objects.filter(household=household, full_name="Baby Aye").exists()


def test_member_of_another_household_fails_the_row(
    household: Household, other_household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    member = other_household.head_of_household
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", member.unicef_id)],
        header=(*DEFAULT_HEADER, "individual_unicef_id"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["failed"] == [
        (2, f"individual {member.unicef_id} belongs to {other_household.unicef_id}, not {household.unicef_id}")
    ]


def test_unknown_member_fails_the_row(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", "IND-00-0000.0000")],
        header=(*DEFAULT_HEADER, "individual_unicef_id"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org")

    assert result["failed"] == [(2, "individual IND-00-0000.0000 not found in programme MCCT")]


def test_repeated_run_does_not_create_the_individual_twice(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)
    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert result["skipped"][0][0] == 2
    assert Individual.objects.filter(household=household, full_name="Baby Aye").count() == 1


def test_rows_sharing_name_and_birth_date_are_both_created(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby Mya", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Mya", "FEMALE", date(2026, 5, 12), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [2, 3]
    assert Individual.objects.filter(household=household, full_name="Baby Mya").count() == 2


def test_repeated_run_does_not_duplicate_rows_sharing_name_and_birth_date(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby Mya", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Mya", "FEMALE", date(2026, 5, 12), "SON_DAUGHTER"),
        ]
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)
    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert Individual.objects.filter(household=household, full_name="Baby Mya").count() == 2


def test_interrupted_run_creates_only_the_individual_that_is_missing(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    IndividualFactory(
        household=household,
        full_name="Baby Mya",
        birth_date=date(2026, 5, 12),
    )
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby Mya", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            ("MCCT", household.unicef_id, "Baby Mya", "FEMALE", date(2026, 5, 12), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [3]
    assert result["skipped"][0][0] == 2
    assert Individual.objects.filter(household=household, full_name="Baby Mya").count() == 2


def test_existing_individual_is_skipped(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    IndividualFactory(
        household=household,
        full_name="Baby Aye",
        birth_date=date(2026, 5, 12),
    )
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == []
    assert "Baby Aye already in" in result["skipped"][0][1]
    assert not GrievanceTicket.objects.exists()


def test_aliased_relationship_column_is_accepted(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")],
        header=("programme_name", "unicef_id", "full_name_i_c", "gender_i_c", "birth_date_i_c", "relationship to hh"),
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.relationship == "SON_DAUGHTER"


def test_header_padded_with_whitespace_still_matches_its_column(
    household: Household, user: User, mother_name_flex_field: FlexibleAttribute, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", "Daw Mya")],
        header=(*DEFAULT_HEADER, " ind_mother_full_name_i_f "),
    )

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.flex_fields["ind_mother_full_name_i_f"] == "Daw Mya"


def test_value_padded_with_whitespace_is_stripped(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "  Baby Aye  ", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert Individual.objects.filter(household=household, full_name="Baby Aye").exists()


def test_text_birth_date_is_read_as_day_first(
    household: Household, user: User, write_workbook: Callable[..., str]
) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", "12/05/2026", "SON_DAUGHTER")])

    bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    individual = Individual.objects.get(household=household, full_name="Baby Aye")
    assert individual.birth_date == date(2026, 5, 12)


def test_computed_age_column_is_ignored(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", 0)],
        header=(*DEFAULT_HEADER, "age"),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [2]
    assert Individual.objects.filter(household=household, full_name="Baby Aye").exists()


@pytest.mark.parametrize(
    ("column", "cell_value"),
    [
        pytest.param("first_registration_date_i_c", date(2020, 1, 1), id="registration_date"),
        pytest.param("photo_i_c", "photo.jpg", id="photo"),
        pytest.param("Remarks", "ok", id="free_text"),
    ],
)
def test_unmatched_column_is_reported_and_ignored(
    household: Household,
    user: User,
    write_workbook: Callable[..., str],
    capsys: pytest.CaptureFixture[str],
    column: str,
    cell_value: Any,
) -> None:
    path = write_workbook(
        [("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER", cell_value)],
        header=(*DEFAULT_HEADER, column),
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [2]
    assert f"ignored columns: ['{column}']" in capsys.readouterr().out
    assert Individual.objects.filter(household=household, full_name="Baby Aye").exists()


def test_blank_rows_are_skipped(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook(
        [
            ("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER"),
            (None, None, None, None, None, None),
            ("MCCT", household.unicef_id, "Baby Two", "FEMALE", date(2026, 6, 1), "SON_DAUGHTER"),
        ]
    )

    result = bulk_add_individuals(path, "myanmar", "loader@unicef.org", apply=True)

    assert result["created"] == [2, 4]
    assert result["failed"] == []


def test_sheet_is_chosen_by_name_when_given(household: Household, user: User, tmp_path: Any) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Disability"
    first.append(list(DEFAULT_HEADER))
    first.append(["MCCT", household.unicef_id, "Wrong Sheet", "MALE", date(2026, 5, 12), "SON_DAUGHTER"])
    second = workbook.create_sheet("MCCT")
    second.append(list(DEFAULT_HEADER))
    second.append(["MCCT", household.unicef_id, "Right Sheet", "MALE", date(2026, 5, 12), "SON_DAUGHTER"])
    path = tmp_path / "two_sheets.xlsx"
    workbook.save(path)

    bulk_add_individuals(str(path), "myanmar", "loader@unicef.org", sheet_name="MCCT", apply=True)

    assert Individual.objects.filter(household=household, full_name="Right Sheet").exists()
    assert not Individual.objects.filter(full_name="Wrong Sheet").exists()


def test_unknown_sheet_name_raises(household: Household, user: User, write_workbook: Callable[..., str]) -> None:
    path = write_workbook([("MCCT", household.unicef_id, "Baby Aye", "MALE", date(2026, 5, 12), "SON_DAUGHTER")])

    with pytest.raises(KeyError, match="Worksheet Wrong does not exist"):
        bulk_add_individuals(path, "myanmar", "loader@unicef.org", sheet_name="Wrong")
