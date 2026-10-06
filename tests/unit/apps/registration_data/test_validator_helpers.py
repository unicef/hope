from collections import defaultdict
from unittest.mock import MagicMock

from openpyxl import Workbook
import pytest

from extras.test_utils.factories.core import BusinessAreaFactory
from extras.test_utils.factories.geo import AreaFactory, AreaTypeFactory, CountryFactory
from extras.test_utils.factories.household import DocumentTypeFactory
from extras.test_utils.factories.program import ProgramFactory
from hope.apps.registration_data.validators import (
    ImportDataInstanceValidator,
    KoboProjectImportDataInstanceValidator,
    UploadXLSXInstanceValidator,
)


@pytest.fixture
def make_header_cells():
    def _factory(values):
        cells = []
        for val in values:
            cell = MagicMock()
            cell.value = val
            cells.append(cell)
        return tuple(cells)

    return _factory


# --- _find_header_column_indices ---


def test_find_header_column_indices_both_present(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    first_row = make_header_cells(["full_name_i_c", "household_id", "relationship_i_c", "age"])
    result = UploadXLSXInstanceValidator._find_header_column_indices(validator, first_row)
    assert result == (1, 2)


def test_find_header_column_indices_only_household_id(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    first_row = make_header_cells(["household_id", "full_name_i_c", "age"])
    result = UploadXLSXInstanceValidator._find_header_column_indices(validator, first_row)
    assert result == (0, None)


def test_find_header_column_indices_only_relationship(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    first_row = make_header_cells(["full_name_i_c", "relationship_i_c", "age"])
    result = UploadXLSXInstanceValidator._find_header_column_indices(validator, first_row)
    assert result == (None, 1)


def test_find_header_column_indices_neither_present(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    first_row = make_header_cells(["full_name_i_c", "age", "gender"])
    result = UploadXLSXInstanceValidator._find_header_column_indices(validator, first_row)
    assert result == (None, None)


def test_find_header_column_indices_empty_row(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    first_row = make_header_cells([])
    result = UploadXLSXInstanceValidator._find_header_column_indices(validator, first_row)
    assert result == (None, None)


# --- _init_doc_identity_dicts ---


def test_init_doc_identity_dicts_returns_two_dicts():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    identities, documents = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert isinstance(identities, dict)
    assert isinstance(documents, dict)


def test_init_doc_identity_dicts_identities_keys():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    identities, _ = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert "unhcr_id_no_i_c" in identities
    assert "scope_id_no_i_c" in identities
    assert identities["unhcr_id_no_i_c"]["partner"] == "UNHCR"
    assert identities["scope_id_no_i_c"]["partner"] == "WFP"


def test_init_doc_identity_dicts_identities_have_empty_lists():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    identities, _ = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert identities["unhcr_id_no_i_c"]["validation_data"] == []
    assert identities["unhcr_id_no_i_c"]["numbers"] == []
    assert identities["unhcr_id_no_i_c"]["issuing_countries"] == []
    assert identities["scope_id_no_i_c"]["validation_data"] == []
    assert identities["scope_id_no_i_c"]["numbers"] == []
    assert identities["scope_id_no_i_c"]["issuing_countries"] == []


def test_init_doc_identity_dicts_documents_keys():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    _, documents = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    expected_doc_keys = {
        "birth_certificate_no_i_c",
        "drivers_license_no_i_c",
        "electoral_card_no_i_c",
        "national_id_no_i_c",
        "national_passport_no_i_c",
        "tax_id_no_i_c",
        "other_id_type_i_c",
        "other_id_no_i_c",
    }
    assert set(documents.keys()) == expected_doc_keys


def test_init_doc_identity_dicts_documents_types():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    _, documents = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert documents["birth_certificate_no_i_c"]["type"] == "BIRTH_CERTIFICATE"
    assert documents["national_passport_no_i_c"]["type"] == "NATIONAL_PASSPORT"
    assert documents["other_id_type_i_c"]["type"] == "OTHER"


def test_init_doc_identity_dicts_other_id_type_has_names_list():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    _, documents = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert "names" in documents["other_id_type_i_c"]
    assert documents["other_id_type_i_c"]["names"] == []


def test_init_doc_identity_dicts_other_id_no_is_empty():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    _, documents = UploadXLSXInstanceValidator._init_doc_identity_dicts(validator)
    assert documents["other_id_no_i_c"] == {}


# --- _validate_admin_column ---


def test_validate_admin_column_non_admin_header_returns_none():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "full_name_i_c", "some_value", 3, [])
    assert result is None


def test_validate_admin_column_admin1_with_value_appends_and_returns_none():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    admin_tuples = []
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "admin1_h_c", "AF01", 3, admin_tuples)
    assert result is None
    assert (3, "admin1_h_c", "AF01") in admin_tuples


def test_validate_admin_column_admin2_with_value_appends_and_returns_none():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    admin_tuples = []
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "admin2_h_c", "AF0101", 5, admin_tuples)
    assert result is None
    assert (5, "admin2_h_c", "AF0101") in admin_tuples


def test_validate_admin_column_admin1_empty_returns_error():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "admin1_h_c", None, 3, [])
    assert result is not None
    assert result["row_number"] == 3
    assert result["header"] == "admin1_h_c"
    assert "cannot be null" in result["message"]


def test_validate_admin_column_admin2_empty_returns_error():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "admin2_h_c", "", 7, [])
    assert result is not None
    assert result["row_number"] == 7
    assert result["header"] == "admin2_h_c"


def test_validate_admin_column_admin3_empty_returns_none():
    """admin3 is not required, so empty value should not produce an error."""
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "admin3_h_c", None, 3, [])
    assert result is None


def test_validate_admin_column_pp_admin3_empty_returns_none():
    """pp_admin3_i_c is also not required."""
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "pp_admin3_i_c", None, 3, [])
    assert result is None


def test_validate_admin_column_pp_admin1_with_value_appends():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    admin_tuples = []
    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "pp_admin1_i_c", "CODE1", 2, admin_tuples)
    assert result is None
    assert (2, "pp_admin1_i_c", "CODE1") in admin_tuples


def test_validate_admin_column_pp_admin4_with_value_appends():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    admin_tuples = []

    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "pp_admin4_i_c", "CODE4", 4, admin_tuples)

    assert result is None
    assert admin_tuples == [(4, "pp_admin4_i_c", "CODE4")]


def test_validate_admin_column_pp_admin4_empty_is_allowed():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL

    result = UploadXLSXInstanceValidator._validate_admin_column(validator, "pp_admin4_i_c", None, 4, [])

    assert result is None


@pytest.fixture
def people_admin4_row():
    sheet = Workbook().active
    sheet["A1"] = "pp_admin4_i_c"
    sheet["A4"] = "MISSING"
    return (sheet["A4"],), (sheet["A1"],)


def test_validate_row_cells_collects_people_area_outside_individual_fields(people_admin4_row):
    validator = object.__new__(UploadXLSXInstanceValidator)
    validator.country_columns = {}
    admin_area_code_tuples = []

    errors = validator._validate_row_cells(*people_admin4_row, {}, admin_area_code_tuples)

    assert errors == []
    assert admin_area_code_tuples == [(4, "pp_admin4_i_c", "MISSING")]


@pytest.fixture
def people_row_with_unlabeled_columns():
    sheet = Workbook().active
    sheet["B1"] = 123
    sheet["C1"] = "pp_country_i_c"
    sheet["D1"] = "pp_birth_certificate_issuer_i_c"
    sheet["E1"] = "unrelated_issuer_i_c"
    sheet["F1"] = "pp_country_origin_i_c"
    sheet["A3"] = "extra under blank header"
    sheet["B3"] = "extra under numeric header"
    sheet["C3"] = "AFG"
    sheet["D3"] = "AF"
    sheet["E3"] = "ZZ"
    sheet["F3"] = "ZZZ"
    return sheet[3], sheet[1]


@pytest.fixture
def people_country_validator():
    validator = object.__new__(UploadXLSXInstanceValidator)
    validator.sheet_title = "People"
    validator.country_columns = {
        "pp_country_i_c": "iso_code3",
        "pp_birth_certificate_issuer_i_c": "iso_code2",
        "pp_unhcr_id_issuer_i_c": "iso_code2",
        "pp_country_origin_i_c": "iso_code3",
    }
    validator.all_fields = {}
    return validator


def test_country_cell_scan_skips_unrecognized_headers(people_row_with_unlabeled_columns, people_country_validator):
    row, headers = people_row_with_unlabeled_columns

    cells = people_country_validator._country_cells_in_row(row, headers)

    assert [(cell.coordinate, header) for cell, header in cells] == [
        ("C3", "pp_country_i_c"),
        ("D3", "pp_birth_certificate_issuer_i_c"),
        ("F3", "pp_country_origin_i_c"),
    ]


@pytest.fixture
def program_with_custom_document_type():
    DocumentTypeFactory(key="custom_document")
    return ProgramFactory()


@pytest.fixture
def custom_document_country_row():
    sheet = Workbook().active
    sheet.title = "Individuals"
    sheet["A1"] = "custom_document_issuer_i_c"
    sheet["A3"] = "ZZ"
    return sheet


@pytest.fixture
def duplicate_custom_document_issuer_row():
    sheet = Workbook().active
    sheet.title = "Individuals"
    sheet["A1"] = "custom_document_issuer_i_c"
    sheet["B1"] = "custom_document_issuer_i_c"
    sheet["A3"] = "ZZ"
    sheet["B3"] = "YY"
    return sheet


@pytest.mark.django_db
def test_validator_collects_configured_document_issuer(program_with_custom_document_type, custom_document_country_row):
    validator = UploadXLSXInstanceValidator(program_with_custom_document_type)

    validator.rows_validator(custom_document_country_row)

    assert validator.errors == [
        {
            "row_number": 3,
            "header": "custom_document_issuer_i_c",
            "message": "Sheet: 'Individuals': Issuing country 'ZZ' does not exist",
        }
    ]


@pytest.mark.django_db
def test_validator_reports_each_duplicate_issuer_column(
    program_with_custom_document_type, duplicate_custom_document_issuer_row
):
    validator = UploadXLSXInstanceValidator(program_with_custom_document_type)

    validator.rows_validator(duplicate_custom_document_issuer_row)

    assert validator.errors == [
        {
            "row_number": 3,
            "header": "custom_document_issuer_i_c",
            "message": "Sheet: 'Individuals': Issuing country 'ZZ' does not exist",
        },
        {
            "row_number": 3,
            "header": "custom_document_issuer_i_c",
            "message": "Sheet: 'Individuals': Issuing country 'YY' does not exist",
        },
    ]


@pytest.fixture(params=["", 0])
def custom_document_invalid_issuer_value(request):
    sheet = Workbook().active
    sheet.title = "Individuals"
    sheet["A1"] = "custom_document_issuer_i_c"
    sheet["B1"] = "custom_document_no_i_c"
    sheet["A3"] = request.param
    sheet["B3"] = "ABC123"
    return sheet, request.param


@pytest.mark.django_db
def test_validator_reports_empty_or_numeric_custom_document_issuer(
    program_with_custom_document_type, custom_document_invalid_issuer_value
):
    sheet, value = custom_document_invalid_issuer_value
    validator = UploadXLSXInstanceValidator(program_with_custom_document_type)

    validator.rows_validator(sheet)

    assert validator.errors == [
        {
            "row_number": 3,
            "header": "custom_document_issuer_i_c",
            "message": f"Sheet: 'Individuals': Issuing country {value!r} does not exist",
        }
    ]


@pytest.fixture
def afghanistan_country():
    return CountryFactory(iso_code2="AF", iso_code3="AFG")


@pytest.fixture
def afghanistan_business_area(afghanistan_country):
    business_area = BusinessAreaFactory(slug="afghanistan")
    business_area.countries.add(afghanistan_country)
    return business_area


@pytest.mark.django_db
def test_validate_admin_areas_reports_missing_people_area(afghanistan_business_area, django_assert_num_queries):
    validator = object.__new__(UploadXLSXInstanceValidator)
    validator.sheet_title = "People"

    with django_assert_num_queries(2):
        errors = validator.validate_admin_areas([(4, "pp_admin4_i_c", "MISSING")], afghanistan_business_area.slug)

    assert errors == [
        {
            "row_number": 4,
            "header": "pp_admin4_i_c",
            "message": "Sheet: 'People': Area with code: MISSING does not exist",
        }
    ]


@pytest.fixture
def mixed_admin_area_codes(afghanistan_country):
    afghanistan_area_type = AreaTypeFactory(country=afghanistan_country)
    AreaFactory(p_code="AF01", area_type=afghanistan_area_type)
    AreaFactory(p_code="AF02", area_type=afghanistan_area_type)
    other_country = CountryFactory(iso_code2="PK", iso_code3="PAK")
    AreaFactory(p_code="PK01", area_type=AreaTypeFactory(country=other_country))
    return [
        (3, "pp_admin1_i_c", "AF01"),
        (3, "pp_admin2_i_c", "AF02"),
        (4, "pp_admin1_i_c", "PK01"),
        (5, "pp_admin1_i_c", "MISSING"),
        (6, "pp_admin2_i_c", "MISSING"),
    ]


@pytest.mark.django_db
def test_validate_admin_areas_batches_queries_and_preserves_cell_errors(
    afghanistan_business_area, mixed_admin_area_codes, django_assert_num_queries
):
    validator = object.__new__(UploadXLSXInstanceValidator)
    validator.sheet_title = "People"

    with django_assert_num_queries(3):
        errors = validator.validate_admin_areas(mixed_admin_area_codes, afghanistan_business_area.slug)

    assert errors == [
        {
            "row_number": 4,
            "header": "pp_admin1_i_c",
            "message": "Sheet: 'People': Admin Area: PK01 unavailable in Business Area: afghanistan",
        },
        {
            "row_number": 5,
            "header": "pp_admin1_i_c",
            "message": "Sheet: 'People': Area with code: MISSING does not exist",
        },
        {
            "row_number": 6,
            "header": "pp_admin2_i_c",
            "message": "Sheet: 'People': Area with code: MISSING does not exist",
        },
    ]


@pytest.fixture
def existing_issuer_country_cells():
    sheet = Workbook().active
    sheet["A4"] = "AFG"
    sheet["B5"] = "AF"
    return [(sheet["A4"], "pp_birth_certificate_issuer_i_c"), (sheet["B5"], "pp_unhcr_id_issuer_i_c")]


@pytest.fixture
def missing_issuer_country_cells():
    sheet = Workbook().active
    sheet["A4"] = "AFG"
    sheet["B5"] = "ZZ"
    return [(sheet["A4"], "pp_birth_certificate_issuer_i_c"), (sheet["B5"], "pp_unhcr_id_issuer_i_c")]


@pytest.mark.django_db
def test_validate_countries_accepts_existing_issuing_country(
    afghanistan_country, people_country_validator, existing_issuer_country_cells, django_assert_num_queries
):

    with django_assert_num_queries(1):
        errors = people_country_validator.validate_countries(existing_issuer_country_cells, {})

    assert errors == []


@pytest.mark.django_db
def test_validate_countries_reports_missing_issuing_country(
    afghanistan_country, people_country_validator, missing_issuer_country_cells, django_assert_num_queries
):

    with django_assert_num_queries(1):
        errors = people_country_validator.validate_countries(missing_issuer_country_cells, {})

    assert errors == [
        {
            "row_number": 5,
            "header": "pp_unhcr_id_issuer_i_c",
            "message": "Sheet: 'People': Issuing country 'ZZ' does not exist",
        }
    ]


# --- _validate_head_of_household (on UploadXLSXInstanceValidator) ---


def test_validate_head_of_household_all_valid():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.head_of_household_count = defaultdict(int, {"HH1": 1, "HH2": 1})
    result = UploadXLSXInstanceValidator._validate_head_of_household(validator)
    assert result == []


def test_validate_head_of_household_missing_head():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.head_of_household_count = defaultdict(int, {"HH1": 0})
    result = UploadXLSXInstanceValidator._validate_head_of_household(validator)
    assert len(result) == 1
    assert "has to have a head of household" in result[0]["message"]
    assert result[0]["header"] == "relationship_i_c"
    assert result[0]["row_number"] == 0


def test_validate_head_of_household_multiple_heads():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.head_of_household_count = defaultdict(int, {"HH1": 3})
    result = UploadXLSXInstanceValidator._validate_head_of_household(validator)
    assert len(result) == 1
    assert "multiple head of households" in result[0]["message"]


def test_validate_head_of_household_mixed_errors():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.head_of_household_count = defaultdict(int, {"HH1": 0, "HH2": 2, "HH3": 1})
    result = UploadXLSXInstanceValidator._validate_head_of_household(validator)
    assert len(result) == 2
    messages = [r["message"] for r in result]
    assert any("has to have a head of household" in m for m in messages)
    assert any("multiple head of households" in m for m in messages)


def test_validate_head_of_household_empty_dict():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.head_of_household_count = defaultdict(int)
    result = UploadXLSXInstanceValidator._validate_head_of_household(validator)
    assert result == []


# --- _validate_household_roles (on KoboProjectImportDataInstanceValidator) ---


def test_validate_household_roles_all_valid():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 1, 1)
    assert result == []


def test_validate_household_roles_no_head():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 0, 1, 0)
    assert len(result) == 1
    assert result[0]["header"] == "relationship_i_c"
    assert "has to have a head of household" in result[0]["message"]


def test_validate_household_roles_multiple_heads():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 2, 1, 0)
    assert len(result) == 1
    assert "Only one person can be a head" in result[0]["message"]


def test_validate_household_roles_no_primary_collector():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 0, 0)
    assert len(result) == 1
    assert result[0]["header"] == "role_i_c"
    assert "must have a primary collector" in result[0]["message"]


def test_validate_household_roles_multiple_primary_collectors():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 2, 0)
    assert len(result) == 1
    assert "Only one person can be a primary collector" in result[0]["message"]


def test_validate_household_roles_multiple_alternate_collectors():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 1, 2)
    assert len(result) == 1
    assert "Only one person can be a alternate collector" in result[0]["message"]


def test_validate_household_roles_all_errors():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 0, 0, 2)
    assert len(result) == 3
    headers = [e["header"] for e in result]
    assert "relationship_i_c" in headers
    assert "role_i_c" in headers


def test_validate_household_roles_zero_alternate_is_ok():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 1, 0)
    assert result == []


def test_validate_household_roles_exactly_one_of_each():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    result = KoboProjectImportDataInstanceValidator._validate_household_roles(validator, 1, 1, 1)
    assert result == []


# --- _process_row_household_data ---


def test_process_row_household_data_no_hh_idx(make_header_cells):
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = []
    validator.head_of_household_count = defaultdict(int)
    row = make_header_cells(["some_value"])
    result = UploadXLSXInstanceValidator._process_row_household_data(validator, row, None, None, "Individuals")
    assert result is None


def test_process_row_household_data_households_sheet():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = []
    validator.head_of_household_count = defaultdict(int)
    cell = MagicMock()
    cell.value = "HH1"
    row = (cell,)
    result = UploadXLSXInstanceValidator._process_row_household_data(validator, row, 0, None, "Households")
    assert result == "HH1"
    assert "HH1" in validator.household_ids
    assert validator.head_of_household_count["HH1"] == 0


def test_process_row_household_data_head_increments():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = []
    validator.head_of_household_count = defaultdict(int)
    hh_cell = MagicMock()
    hh_cell.value = "HH1"
    rel_cell = MagicMock()
    rel_cell.value = "HEAD"
    row = (hh_cell, rel_cell)
    result = UploadXLSXInstanceValidator._process_row_household_data(validator, row, 0, 1, "Individuals")
    assert result == "HH1"
    assert validator.head_of_household_count["HH1"] == 1


def test_process_row_household_data_float_to_int():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = []
    validator.head_of_household_count = defaultdict(int)
    cell = MagicMock()
    cell.value = 1.0
    row = (cell,)
    result = UploadXLSXInstanceValidator._process_row_household_data(validator, row, 0, None, "Households")
    assert result == 1
    assert isinstance(result, int)
    assert 1 in validator.household_ids


# --- _validate_field_type ---


def test_validate_field_type_valid():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    fn = MagicMock(return_value=True)
    validator.switch_dict = {"STRING": fn}
    validator.sheet_title = "Individuals"
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    cell = MagicMock()
    cell.row = 3
    result = UploadXLSXInstanceValidator._validate_field_type(
        validator, "some_value", "name_field", cell, "STRING", False
    )
    assert result is None


def test_validate_field_type_invalid():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    fn = MagicMock(return_value=False)
    validator.switch_dict = {"STRING": fn}
    validator.sheet_title = "Individuals"
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    cell = MagicMock()
    cell.row = 5
    result = UploadXLSXInstanceValidator._validate_field_type(
        validator, "bad_value", "name_field", cell, "STRING", False
    )
    assert result is not None
    assert result["row_number"] == 5
    assert result["header"] == "name_field"
    assert "Unexpected value" in result["message"]


def test_validate_field_type_admin_column_skipped():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    fn = MagicMock(return_value=False)
    validator.switch_dict = {"STRING": fn}
    validator.sheet_title = "Households"
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    cell = MagicMock()
    cell.row = 3
    # admin1_h_c is in ADMIN_COLUMNS_ALL, so even with fn returning False, no error is returned
    result = UploadXLSXInstanceValidator._validate_field_type(
        validator, "bad_value", "admin1_h_c", cell, "STRING", False
    )
    assert result is None


# --- _validate_row_household_reference ---


def test_validate_row_household_reference_missing():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = ["HH1", "HH2"]
    result = UploadXLSXInstanceValidator._validate_row_household_reference(validator, "HH999", 10)
    assert result is not None
    assert result["row_number"] == 10
    assert result["header"] == "relationship_i_c"
    assert "no household" in result["message"].lower() or "There is no household" in result["message"]


def test_validate_row_household_reference_exists():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = ["HH1", "HH2"]
    result = UploadXLSXInstanceValidator._validate_row_household_reference(validator, "HH1", 10)
    assert result is None


# --- _run_document_identity_validation ---


def test_run_doc_identity_validation_individuals_calls_validators():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.documents_validator = MagicMock(return_value=[{"error": "doc_error"}])
    validator.identity_validator = MagicMock(return_value=[{"error": "ident_error"}])
    docs = {"some_doc": {}}
    idents = {"some_ident": {}}
    validator._documents_numbers = docs
    validator._identities_numbers = idents
    invalid_doc_rows, invalid_ident_rows = UploadXLSXInstanceValidator._run_document_identity_validation(
        validator, "Individuals"
    )
    validator.documents_validator.assert_called_once_with(docs)
    validator.identity_validator.assert_called_once_with(idents)
    assert invalid_doc_rows == [{"error": "doc_error"}]
    assert invalid_ident_rows == [{"error": "ident_error"}]


def test_run_doc_identity_validation_households_returns_empty():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.documents_validator = MagicMock()
    validator.identity_validator = MagicMock()
    validator._documents_numbers = {}
    validator._identities_numbers = {}
    invalid_doc_rows, invalid_ident_rows = UploadXLSXInstanceValidator._run_document_identity_validation(
        validator, "Households"
    )
    validator.documents_validator.assert_not_called()
    validator.identity_validator.assert_not_called()
    assert invalid_doc_rows == []
    assert invalid_ident_rows == []


# --- _accumulate_doc_identity_validation_data ---


def test_accumulate_data_adds_to_lists():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    # Use the real DOCUMENTS_ISSUING_COUNTRIES_MAPPING from ImportDataInstanceValidator
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING

    # Build identities_numbers and documents_numbers with empty validation_data lists
    # The mapping values are the keys we need in these dicts
    identities_numbers = {
        "unhcr_id_no_i_c": {"validation_data": []},
        "scope_id_no_i_c": {"validation_data": []},
    }
    documents_numbers = {
        "birth_certificate_no_i_c": {"validation_data": []},
        "drivers_license_no_i_c": {"validation_data": []},
        "electoral_card_no_i_c": {"validation_data": []},
        "national_id_no_i_c": {"validation_data": []},
        "national_passport_no_i_c": {"validation_data": []},
        "tax_id_no_i_c": {"validation_data": []},
        "other_id_type_i_c": {"validation_data": []},
    }
    validator._identities_numbers = identities_numbers
    validator._documents_numbers = documents_numbers

    # Create a mock row where row[0].row returns a row number
    first_cell = MagicMock()
    first_cell.row = 7
    row = (first_cell,)

    UploadXLSXInstanceValidator._accumulate_doc_identity_validation_data(validator, row)

    # Each value in the mapping should have had a validation_data entry appended
    assert {"row_number": 7} in identities_numbers["unhcr_id_no_i_c"]["validation_data"]
    assert {"row_number": 7} in identities_numbers["scope_id_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["birth_certificate_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["drivers_license_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["electoral_card_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["national_id_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["national_passport_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["tax_id_no_i_c"]["validation_data"]
    assert {"row_number": 7} in documents_numbers["other_id_type_i_c"]["validation_data"]


# --- _process_document_number ---


def test_process_document_number_regular_doc_appends_number():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {
        "birth_certificate_no_i_c": {"numbers": [], "names": [], "issuing_countries": []},
    }
    identities_numbers = {}
    UploadXLSXInstanceValidator._process_document_number(
        validator, "birth_certificate_no_i_c", "DOC-123", documents_numbers, identities_numbers
    )
    assert "DOC-123" in documents_numbers["birth_certificate_no_i_c"]["numbers"]


def test_process_document_number_other_id_type_appends_name():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {
        "other_id_type_i_c": {"numbers": [], "names": [], "issuing_countries": []},
    }
    identities_numbers = {}
    UploadXLSXInstanceValidator._process_document_number(
        validator, "other_id_type_i_c", "Custom ID", documents_numbers, identities_numbers
    )
    assert "Custom ID" in documents_numbers["other_id_type_i_c"]["names"]


def test_process_document_number_other_id_no_appends_to_other_id_type_numbers():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {
        "other_id_type_i_c": {"numbers": [], "names": [], "issuing_countries": []},
        "other_id_no_i_c": {},
    }
    identities_numbers = {}
    UploadXLSXInstanceValidator._process_document_number(
        validator, "other_id_no_i_c", "99", documents_numbers, identities_numbers
    )
    assert "99" in documents_numbers["other_id_type_i_c"]["numbers"]


def test_process_document_number_issuing_country_in_documents():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {
        "birth_certificate_no_i_c": {"numbers": [], "names": [], "issuing_countries": []},
    }
    identities_numbers = {}
    UploadXLSXInstanceValidator._process_document_number(
        validator, "birth_certificate_issuer_i_c", "AFG", documents_numbers, identities_numbers
    )
    assert "AFG" in documents_numbers["birth_certificate_no_i_c"]["issuing_countries"]


def test_process_document_number_issuing_country_in_identities():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {}
    identities_numbers = {
        "unhcr_id_no_i_c": {"numbers": [], "issuing_countries": []},
    }
    UploadXLSXInstanceValidator._process_document_number(
        validator, "unhcr_id_issuer_i_c", "SYR", documents_numbers, identities_numbers
    )
    assert "SYR" in identities_numbers["unhcr_id_no_i_c"]["issuing_countries"]


def test_process_document_number_identity_appends_number():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {}
    identities_numbers = {
        "unhcr_id_no_i_c": {"numbers": [], "issuing_countries": []},
    }
    UploadXLSXInstanceValidator._process_document_number(
        validator, "unhcr_id_no_i_c", "UNHCR-001", documents_numbers, identities_numbers
    )
    assert "UNHCR-001" in identities_numbers["unhcr_id_no_i_c"]["numbers"]


def test_process_document_number_none_value_appends_none():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {
        "birth_certificate_no_i_c": {"numbers": [], "names": [], "issuing_countries": []},
    }
    identities_numbers = {}
    UploadXLSXInstanceValidator._process_document_number(
        validator, "birth_certificate_no_i_c", None, documents_numbers, identities_numbers
    )
    assert None in documents_numbers["birth_certificate_no_i_c"]["numbers"]


def test_process_document_number_not_in_any_dict():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING = UploadXLSXInstanceValidator.DOCUMENTS_ISSUING_COUNTRIES_MAPPING
    documents_numbers = {}
    identities_numbers = {}
    # Should not raise — just a no-op
    UploadXLSXInstanceValidator._process_document_number(
        validator, "unknown_field", "value", documents_numbers, identities_numbers
    )


# --- _validate_field_type (household_id_can_be_empty=True) ---


def test_validate_field_type_invalid_but_household_id_can_be_empty():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    fn = MagicMock(return_value=False)
    validator.switch_dict = {"STRING": fn}
    validator.sheet_title = "Individuals"
    validator.ADMIN_COLUMNS_ALL = UploadXLSXInstanceValidator.ADMIN_COLUMNS_ALL
    cell = MagicMock()
    cell.row = 5
    result = UploadXLSXInstanceValidator._validate_field_type(
        validator, "bad_value", "name_field", cell, "STRING", True
    )
    assert result is None


# --- _process_row_household_data (float non-integer) ---


def test_process_row_household_data_float_non_integer():
    validator = MagicMock(spec=UploadXLSXInstanceValidator)
    validator.household_ids = []
    validator.head_of_household_count = defaultdict(int)
    cell = MagicMock()
    cell.value = 1.5
    row = (cell,)
    result = UploadXLSXInstanceValidator._process_row_household_data(validator, row, 0, None, "Individuals")
    # 1.5 is float but NOT .is_integer() → NOT converted to int, stays as 1.5
    assert result == 1.5
    assert isinstance(result, float)


# --- _count_individual_roles ---


def test_count_individual_roles_head():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    h, p, a = KoboProjectImportDataInstanceValidator._count_individual_roles(
        validator, "relationship_i_c", "head", 0, 0, 0
    )
    assert (h, p, a) == (1, 0, 0)


def test_count_individual_roles_primary():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    h, p, a = KoboProjectImportDataInstanceValidator._count_individual_roles(validator, "role_i_c", "primary", 0, 0, 0)
    assert (h, p, a) == (0, 1, 0)


def test_count_individual_roles_alternate():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    h, p, a = KoboProjectImportDataInstanceValidator._count_individual_roles(
        validator, "role_i_c", "alternate", 0, 0, 0
    )
    assert (h, p, a) == (0, 0, 1)


def test_count_individual_roles_unrelated_field():
    validator = MagicMock(spec=KoboProjectImportDataInstanceValidator)
    h, p, a = KoboProjectImportDataInstanceValidator._count_individual_roles(
        validator, "full_name_i_c", "John", 5, 3, 1
    )
    assert (h, p, a) == (5, 3, 1)  # unchanged


# --- _make_doc_error ---


def test_make_doc_error_with_xlsx():
    result = ImportDataInstanceValidator._make_doc_error("doc_key", "some error", 5, is_xlsx=True)
    assert result == {"header": "doc_key", "message": "some error", "row_number": 5}


def test_make_doc_error_without_xlsx():
    result = ImportDataInstanceValidator._make_doc_error("doc_key", "some error", 5, is_xlsx=False)
    assert result == {"header": "doc_key", "message": "some error"}
    assert "row_number" not in result


# --- _cell_has_value ---


def test_cell_has_value_none():
    cell = MagicMock()
    cell.value = None
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is False


def test_cell_has_value_empty_string():
    cell = MagicMock()
    cell.value = "   "
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is False


def test_cell_has_value_non_empty_string():
    cell = MagicMock()
    cell.value = "hello"
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is True


def test_cell_has_value_number():
    cell = MagicMock()
    cell.value = 42
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is True


def test_cell_has_value_zero():
    cell = MagicMock()
    cell.value = 0
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is True


def test_cell_has_value_false():
    cell = MagicMock()
    cell.value = False
    assert UploadXLSXInstanceValidator._cell_has_value(cell) is True


# --- _validate_other_id_docs ---


def test_validate_other_id_docs_name_missing_when_number_present():
    validator = MagicMock(spec=ImportDataInstanceValidator)
    validator._make_doc_error = ImportDataInstanceValidator._make_doc_error
    values = {
        "names": [None],
        "numbers": ["12345"],
        "validation_data": [{"row_number": 3}],
    }
    errors = ImportDataInstanceValidator._validate_other_id_docs(
        validator, values, "other_id_type_i_c", ["AFG"], is_xlsx=True
    )
    assert len(errors) == 1
    assert "Name for" in errors[0]["message"]
    assert errors[0]["row_number"] == 3


def test_validate_other_id_docs_number_missing_when_name_present():
    validator = MagicMock(spec=ImportDataInstanceValidator)
    validator._make_doc_error = ImportDataInstanceValidator._make_doc_error
    values = {
        "names": ["Custom ID"],
        "numbers": [None],
        "validation_data": [{"row_number": 5}],
    }
    errors = ImportDataInstanceValidator._validate_other_id_docs(
        validator, values, "other_id_type_i_c", ["AFG"], is_xlsx=True
    )
    assert len(errors) == 1
    assert "Number for" in errors[0]["message"]


def test_validate_other_id_docs_issuing_country_missing():
    validator = MagicMock(spec=ImportDataInstanceValidator)
    validator._make_doc_error = ImportDataInstanceValidator._make_doc_error
    values = {
        "names": ["Custom ID"],
        "numbers": ["12345"],
        "validation_data": [{"row_number": 7}],
    }
    errors = ImportDataInstanceValidator._validate_other_id_docs(
        validator, values, "other_id_type_i_c", [None], is_xlsx=True
    )
    assert len(errors) == 1
    assert "Issuing country" in errors[0]["message"]


def test_validate_other_id_docs_all_valid():
    validator = MagicMock(spec=ImportDataInstanceValidator)
    validator._make_doc_error = ImportDataInstanceValidator._make_doc_error
    values = {
        "names": ["Custom ID"],
        "numbers": ["12345"],
        "validation_data": [{"row_number": 1}],
    }
    errors = ImportDataInstanceValidator._validate_other_id_docs(
        validator, values, "other_id_type_i_c", ["AFG"], is_xlsx=True
    )
    assert len(errors) == 0


def test_validate_other_id_docs_without_xlsx():
    validator = MagicMock(spec=ImportDataInstanceValidator)
    validator._make_doc_error = ImportDataInstanceValidator._make_doc_error
    values = {
        "names": [None],
        "numbers": ["12345"],
        "validation_data": [{"row_number": 3}],
    }
    errors = ImportDataInstanceValidator._validate_other_id_docs(
        validator, values, "other_id_type_i_c", ["AFG"], is_xlsx=False
    )
    assert len(errors) == 1
    assert "row_number" not in errors[0]
