from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
import pytest

from hope.apps.payment.xlsx.utils import get_last_data_row


def _save_and_load(wb: openpyxl.Workbook) -> Worksheet:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return openpyxl.load_workbook(buffer, data_only=True).active


@pytest.fixture
def worksheet(request) -> Worksheet:
    wb = openpyxl.Workbook()
    for row in request.param:
        wb.active.append(row)
    return _save_and_load(wb)


@pytest.fixture
def worksheet_with_styled_empty_trailing_cell() -> Worksheet:
    wb = openpyxl.Workbook()
    wb.active.append(["payment_id"])
    wb.active.append(["P-1"])
    wb.active["F1048576"].font = Font(bold=True)  # without styling openpyxl doesn't save the empty cell
    return _save_and_load(wb)


@pytest.mark.parametrize(
    ("worksheet", "expected"),
    [
        ([["payment_id"], ["P-1"], ["P-2"]], 3),
        ([["payment_id"], ["P-1"], [None], ["P-2"]], 4),
        ([["payment_id"]], 1),
        ([], 1),
    ],
    ids=["contiguous", "blank_row_in_the_middle", "header_only", "empty_sheet"],
    indirect=["worksheet"],
)
def test_get_last_data_row_returns_last_row_with_a_value(worksheet, expected):
    result = get_last_data_row(worksheet)

    assert result == expected


def test_get_last_data_row_ignores_styled_empty_trailing_cell(worksheet_with_styled_empty_trailing_cell):
    result = get_last_data_row(worksheet_with_styled_empty_trailing_cell)

    assert result == 2


def test_get_last_data_row_does_not_create_cells(worksheet_with_styled_empty_trailing_cell):
    cells_before = len(worksheet_with_styled_empty_trailing_cell._cells)

    get_last_data_row(worksheet_with_styled_empty_trailing_cell)

    assert len(worksheet_with_styled_empty_trailing_cell._cells) == cells_before
