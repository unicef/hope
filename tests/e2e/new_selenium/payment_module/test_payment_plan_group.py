from decimal import Decimal
from unittest import mock

from django.utils import timezone
import openpyxl
import pytest

from e2e.new_selenium.conftest import grant_permission
from extras.test_utils.factories import (
    DeliveryMechanismFactory,
    FileTempFactory,
    FinancialServiceProviderFactory,
    FinancialServiceProviderXlsxTemplateFactory,
    FspXlsxTemplatePerDeliveryMechanismFactory,
    PaymentFactory,
    PaymentHouseholdSnapshotFactory,
    PaymentPlanFactory,
    PaymentPlanGroupFactory,
    PaymentPlanSplitFactory,
)
from extras.test_utils.selenium import HopeTestBrowser
from hope.apps.account.permissions import Permissions
from hope.apps.payment.flows import PaymentPlanFlow
from hope.apps.payment.xlsx.xlsx_payment_plan_group_delivery_export_service import (
    XlsxPaymentPlanGroupDeliveryExportService,
)
from hope.models import (
    BusinessArea,
    DeliveryMechanism,
    FinancialServiceProvider,
    FinancialServiceProviderXlsxTemplate,
    Payment,
    PaymentPlan,
    PaymentPlanGroup,
    ProgramCycle,
    User,
)

pytestmark = pytest.mark.django_db()


@pytest.fixture
def plan_with_follow_up(program_cycle: ProgramCycle) -> tuple[PaymentPlan, PaymentPlan]:
    """A source Payment Plan plus a follow-up child plan pointing at it.

    The eye icon / Linked Payment Plans modal renders only when a plan's follow_ups
    (child_plans with plan_type FOLLOW_UP) is non-empty, so the source plan is what
    shows the eye icon and the follow-up is the record listed inside the modal.
    """
    ba = program_cycle.program.business_area
    source = PaymentPlanFactory(
        program_cycle=program_cycle,
        business_area=ba,
        status=PaymentPlan.Status.FINISHED,
        plan_type=PaymentPlan.PlanType.REGULAR,
    )
    follow_up = PaymentPlanFactory(
        program_cycle=program_cycle,
        business_area=ba,
        status=PaymentPlan.Status.OPEN,
        plan_type=PaymentPlan.PlanType.FOLLOW_UP,
        source_payment_plan=source,
    )
    return source, follow_up


def _assert_linked_plans_modal_opens_and_closes(
    browser: HopeTestBrowser,
    source: PaymentPlan,
    follow_up: PaymentPlan,
) -> None:
    """Open the Linked Payment Plans modal via the eye icon, verify it lists the
    linked follow-up plan, and close it — asserting neither click bubbles to the
    ClickableTableRow and navigates to the plan details page (the stopPropagation fix)."""
    browser.wait_for_text(source.unicef_id)
    browser.wait_for_element_clickable('[data-cy="button-eye-linked-plans"]')
    browser.click('[data-cy="button-eye-linked-plans"]')

    browser.wait_for_element_visible('[data-cy="table-cell-linked-payment-plan-id"]')
    browser.wait_for_text(follow_up.unicef_id, '[role="dialog"]')

    # Regression guard: the eye click must not navigate to the plan details page,
    # so the source plan id never appears in the URL while the modal is open.
    assert str(source.id) not in browser.get_current_url()

    # Close the modal; the Close click must also not bubble/navigate.
    browser.click('[data-cy="button-close"]')
    browser.wait_for_element_not_visible('[data-cy="table-cell-linked-payment-plan-id"]')
    assert str(source.id) not in browser.get_current_url()


@pytest.fixture
def group_fsp() -> FinancialServiceProvider:
    return FinancialServiceProviderFactory(
        name="Group Delivery FSP",
        communication_channel=FinancialServiceProvider.COMMUNICATION_CHANNEL_XLSX,
        vision_vendor_number="900900900",
    )


@pytest.fixture
def group_delivery_mechanism() -> DeliveryMechanism:
    return DeliveryMechanismFactory(code="grp-cash", name="Group Cash", payment_gateway_id="grp-cash")


@pytest.fixture
def group_fsp_template(
    group_fsp: FinancialServiceProvider,
    group_delivery_mechanism: DeliveryMechanism,
) -> FinancialServiceProviderXlsxTemplate:
    return FspXlsxTemplatePerDeliveryMechanismFactory(
        financial_service_provider=group_fsp,
        delivery_mechanism=group_delivery_mechanism,
        xlsx_template=FinancialServiceProviderXlsxTemplateFactory(columns=["payment_id", "delivered_quantity"]),
    ).xlsx_template


@pytest.fixture
def exportable_group(
    program_cycle: ProgramCycle,
    group_fsp: FinancialServiceProvider,
    group_delivery_mechanism: DeliveryMechanism,
    group_fsp_template: FinancialServiceProviderXlsxTemplate,
) -> tuple[PaymentPlanGroup, Payment]:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Exportable Group")
    plan = PaymentPlanFactory(
        program_cycle=program_cycle,
        payment_plan_group=group,
        business_area=program_cycle.program.business_area,
        financial_service_provider=group_fsp,
        delivery_mechanism=group_delivery_mechanism,
        status=PaymentPlan.Status.ACCEPTED,
        plan_type=PaymentPlan.PlanType.REGULAR,
    )
    payment = PaymentFactory(
        parent=plan,
        financial_service_provider=group_fsp,
        delivery_type=group_delivery_mechanism,
        program=plan.program,
        entitlement_quantity=Decimal("100.00"),
        entitlement_quantity_usd=Decimal("10.00"),
    )
    PaymentHouseholdSnapshotFactory(payment=payment, snapshot_data={})
    return group, payment


@pytest.fixture
def auth_code_template(business_area: BusinessArea) -> FinancialServiceProviderXlsxTemplate:
    # The export-with-auth-code dialog lists templates whose FSP is allowed in the business
    # area; this one carries the fsp_auth_code column so the dialog's auth-code path is exercised.
    fsp = FinancialServiceProviderFactory(
        name="Auth Code FSP",
        communication_channel=FinancialServiceProvider.COMMUNICATION_CHANNEL_XLSX,
        vision_vendor_number="900900902",
    )
    fsp.allowed_business_areas.add(business_area)
    return FspXlsxTemplatePerDeliveryMechanismFactory(
        financial_service_provider=fsp,
        delivery_mechanism=DeliveryMechanismFactory(code="auth-cash", name="Auth Cash", payment_gateway_id="auth-cash"),
        xlsx_template=FinancialServiceProviderXlsxTemplateFactory(
            name="Auth Code Template",
            columns=["payment_id", "delivered_quantity", "fsp_auth_code"],
        ),
    ).xlsx_template


@pytest.fixture
def reconciliation_file(tmp_path, exportable_group: tuple[PaymentPlanGroup, Payment]) -> str:
    group, payment = exportable_group
    # Build the file from the real export service so its header matches exactly what the
    # import expects, then fill in a delivered_quantity for the single payment row.
    workbook = XlsxPaymentPlanGroupDeliveryExportService(
        group, plan_type=PaymentPlan.PlanType.REGULAR
    ).generate_workbook()
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    delivered_col = headers.index("delivered_quantity") + 1
    worksheet.cell(row=2, column=delivered_col).value = 50
    file_path = tmp_path / "reconciliation.xlsx"
    workbook.save(str(file_path))
    payment.status = Payment.STATUS_SENT_TO_FSP
    payment.status_date = timezone.now()
    payment.save(update_fields=["status", "status_date"])
    return str(file_path)


def _set_delivered_quantities(file_path: str, quantities: dict[str, int | None]) -> None:
    workbook = openpyxl.load_workbook(file_path)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    payment_id_column = headers.index("payment_id")
    delivered_quantity_column = headers.index("delivered_quantity")
    updated_payment_ids: set[str] = set()
    for row in worksheet.iter_rows(min_row=2):
        payment_id = str(row[payment_id_column].value)
        if payment_id in quantities:
            row[delivered_quantity_column].value = quantities[payment_id]
            updated_payment_ids.add(payment_id)
    missing_payment_ids = set(quantities) - updated_payment_ids
    if missing_payment_ids:
        raise AssertionError(f"Payments not found in reconciliation XLSX: {sorted(missing_payment_ids)}")
    workbook.save(file_path)


@pytest.fixture
def finished_group_with_empty_reconciliation_file(
    tmp_path,
    exportable_group: tuple[PaymentPlanGroup, Payment],
) -> tuple[PaymentPlanGroup, Payment, str]:
    group, payment = exportable_group
    payment_plan = payment.parent
    workbook = XlsxPaymentPlanGroupDeliveryExportService(
        group, plan_type=PaymentPlan.PlanType.REGULAR
    ).generate_workbook()
    file_path = tmp_path / "empty_reconciliation.xlsx"
    workbook.save(file_path)

    payment.delivered_quantity = Decimal("50.00")
    payment.delivered_quantity_usd = Decimal("5.00")
    payment.delivery_date = timezone.now()
    payment.status = Payment.STATUS_DISTRIBUTION_PARTIAL
    payment.status_date = timezone.now()
    payment.transaction_reference_id = "ORIGINAL-REFERENCE"
    payment.reason_for_unsuccessful_payment = "Original reason"
    payment.additional_collector_name = "Original collector"
    payment.additional_document_type = "National ID"
    payment.additional_document_number = "ORIGINAL-DOCUMENT"
    payment.transaction_status_blockchain_link = "https://example.com/original-transaction"
    payment.set_extra_fields({"reconciliation_note": "original"})
    payment.set_fsp_extra_fields({"fsp_reference": "keep"})
    payment.save(
        update_fields=[
            "delivered_quantity",
            "delivered_quantity_usd",
            "delivery_date",
            "status",
            "status_date",
            "transaction_reference_id",
            "reason_for_unsuccessful_payment",
            "additional_collector_name",
            "additional_document_type",
            "additional_document_number",
            "transaction_status_blockchain_link",
            "extras",
        ]
    )
    PaymentPlanFlow(payment_plan).status_finished()
    payment_plan.update_money_fields()
    payment_plan.save()
    return group, payment, str(file_path)


@pytest.fixture
def mixed_reconciliation_group(
    program_cycle: ProgramCycle,
    group_fsp: FinancialServiceProvider,
    group_delivery_mechanism: DeliveryMechanism,
    group_fsp_template: FinancialServiceProviderXlsxTemplate,
) -> tuple[PaymentPlanGroup, PaymentPlan, PaymentPlan, Payment, Payment]:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Mixed Reconciliation Group")
    payment_plans = [
        PaymentPlanFactory(
            program_cycle=program_cycle,
            payment_plan_group=group,
            business_area=program_cycle.program.business_area,
            financial_service_provider=group_fsp,
            delivery_mechanism=group_delivery_mechanism,
            status=PaymentPlan.Status.ACCEPTED,
            plan_type=PaymentPlan.PlanType.REGULAR,
        )
        for _ in range(2)
    ]
    payments = [
        PaymentFactory(
            parent=payment_plan,
            financial_service_provider=group_fsp,
            delivery_type=group_delivery_mechanism,
            program=payment_plan.program,
            entitlement_quantity=Decimal("100.00"),
            entitlement_quantity_usd=Decimal("10.00"),
        )
        for payment_plan in payment_plans
    ]
    for payment in payments:
        PaymentHouseholdSnapshotFactory(payment=payment, snapshot_data={})
    return group, payment_plans[0], payment_plans[1], payments[0], payments[1]


@pytest.fixture
def mixed_reconciliation_file(
    tmp_path,
    mixed_reconciliation_group: tuple[PaymentPlanGroup, PaymentPlan, PaymentPlan, Payment, Payment],
) -> str:
    group, _, _, first_payment, second_payment = mixed_reconciliation_group
    workbook = XlsxPaymentPlanGroupDeliveryExportService(
        group, plan_type=PaymentPlan.PlanType.REGULAR
    ).generate_workbook()
    file_path = tmp_path / "mixed_reconciliation.xlsx"
    workbook.save(file_path)
    _set_delivered_quantities(str(file_path), {str(first_payment.unicef_id): 50})

    for payment in (first_payment, second_payment):
        payment.status = Payment.STATUS_SENT_TO_FSP
        payment.status_date = timezone.now()
        payment.save(update_fields=["status", "status_date"])
    return str(file_path)


@pytest.fixture
def sendable_group(program_cycle: ProgramCycle) -> tuple[PaymentPlanGroup, PaymentPlan]:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Sendable Group")
    plan = PaymentPlanFactory(
        program_cycle=program_cycle,
        payment_plan_group=group,
        business_area=program_cycle.program.business_area,
        status=PaymentPlan.Status.ACCEPTED,
        financial_service_provider=FinancialServiceProviderFactory(name="PG FSP", vision_vendor_number="900900901"),
        use_payment_gateway=True,
    )
    PaymentPlanSplitFactory(payment_plan=plan)
    return group, plan


@pytest.fixture
def group_with_exported_batch(program_cycle: ProgramCycle) -> PaymentPlanGroup:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Batch Group")
    PaymentPlanFactory(
        program_cycle=program_cycle,
        payment_plan_group=group,
        business_area=program_cycle.program.business_area,
        status=PaymentPlan.Status.ACCEPTED,
        export_tag=1,
        export_file_delivery=FileTempFactory(),
    )
    return group


@pytest.fixture
def group_with_unexported_batch(program_cycle: ProgramCycle) -> PaymentPlanGroup:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Pending Batch Group")
    PaymentPlanFactory(
        program_cycle=program_cycle,
        payment_plan_group=group,
        business_area=program_cycle.program.business_area,
        status=PaymentPlan.Status.ACCEPTED,
        export_tag=1,
        export_file_delivery=None,
    )
    return group


@pytest.fixture
def busy_group(exportable_group: tuple[PaymentPlanGroup, Payment]) -> PaymentPlanGroup:
    """Exportable group stuck mid-export, so the XLSX buttons render but disabled."""
    group, _ = exportable_group
    group.background_action_status = PaymentPlanGroup.BackgroundActionStatus.XLSX_EXPORTING
    group.save(update_fields=["background_action_status"])
    return group


@pytest.fixture
def group_with_totals(program_cycle: ProgramCycle) -> PaymentPlanGroup:
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Totals Group")
    PaymentPlanFactory(
        program_cycle=program_cycle,
        payment_plan_group=group,
        business_area=program_cycle.program.business_area,
        status=PaymentPlan.Status.ACCEPTED,
        total_entitled_quantity_usd=Decimal("125.50"),
        total_delivered_quantity_usd=Decimal("25.50"),
        total_undelivered_quantity_usd=Decimal("100.00"),
    )
    return group


@pytest.fixture
def multi_plan_type_group(
    program_cycle: ProgramCycle,
    group_fsp: FinancialServiceProvider,
    group_delivery_mechanism: DeliveryMechanism,
    group_fsp_template: FinancialServiceProviderXlsxTemplate,
) -> tuple[PaymentPlanGroup, PaymentPlan, PaymentPlan]:
    """Group with two exportable plan types, so the export dialog shows a plan type select.

    A single exportable plan type renders a locked read-only field instead; two
    ACCEPTED plans with no export_tag make both canExportRegular and canExportTopUp true.
    """
    group = PaymentPlanGroupFactory(cycle=program_cycle, name="Multi Plan Type Group")

    def build_plan(plan_type: str) -> PaymentPlan:
        plan = PaymentPlanFactory(
            program_cycle=program_cycle,
            payment_plan_group=group,
            business_area=program_cycle.program.business_area,
            financial_service_provider=group_fsp,
            delivery_mechanism=group_delivery_mechanism,
            status=PaymentPlan.Status.ACCEPTED,
            plan_type=plan_type,
        )
        payment = PaymentFactory(
            parent=plan,
            financial_service_provider=group_fsp,
            delivery_type=group_delivery_mechanism,
            program=plan.program,
            entitlement_quantity=Decimal("100.00"),
            entitlement_quantity_usd=Decimal("10.00"),
        )
        PaymentHouseholdSnapshotFactory(payment=payment, snapshot_data={})
        return plan

    return (
        group,
        build_plan(PaymentPlan.PlanType.REGULAR),
        build_plan(PaymentPlan.PlanType.TOP_UP),
    )


def test_create_payment_plan_group(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    program_cycle: ProgramCycle,
) -> None:
    program = program_cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PROGRAMME_CYCLE_VIEW_LIST,
        Permissions.PM_PROGRAMME_CYCLE_VIEW_DETAILS,
        Permissions.PM_PAYMENT_PLAN_GROUP_CREATE,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/program-cycles")

        browser.click('[data-cy="program-cycle-title"] a')
        browser.wait_for_element_clickable('[data-cy="button-create-payment-plan-group"]')
        browser.click('[data-cy="button-create-payment-plan-group"]')

        browser.wait_for_element_visible('input[name="groupName"]')
        browser.type('input[name="groupName"]', "E2E Test Group")
        browser.click('[data-cy="button-create-group-submit"]')
        browser.wait_for_text("Payment Plan Group created")

        browser.click('[data-cy="nav-Payment Module"]')
        browser.wait_for_element_clickable('a[data-cy="nav-Groups"]')
        browser.click('a[data-cy="nav-Groups"]')

        browser.wait_for_text("E2E Test Group")


def test_edit_payment_plan_group(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    payment_plan_group: PaymentPlanGroup,
) -> None:
    program = payment_plan_group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_UPDATE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{payment_plan_group.id}")

        browser.wait_for_element_clickable('[data-cy="button-edit-group-name"]')
        browser.click('[data-cy="button-edit-group-name"]')

        browser.wait_for_element_visible('input[name="name"]')
        browser.clear('input[name="name"]')
        browser.type('input[name="name"]', "Updated Group Name")
        browser.click('[data-cy="button-submit"]')

        browser.wait_for_text("Group name updated")
        browser.assert_text("Updated Group Name")


def test_delete_payment_plan_group(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    payment_plan_group: PaymentPlanGroup,
    second_payment_plan_group: PaymentPlanGroup,
) -> None:
    program = payment_plan_group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PROGRAMME_CYCLE_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_DELETE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{payment_plan_group.id}")

        browser.wait_for_element_clickable('[data-cy="button-delete-group"]')
        browser.click('[data-cy="button-delete-group"]')

        browser.wait_for_text("Are you sure you want to remove this Group?")
        browser.wait_for_element_clickable('[role="dialog"] [data-cy="button-submit"]')
        browser.find_element('[role="dialog"] [data-cy="button-submit"]').click()
        browser.wait_for_text("Group Deleted", timeout=20)

        browser.wait_for_text(second_payment_plan_group.name)
        browser.assert_text_not_visible(payment_plan_group.name)


def test_export_payment_plan_group(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    exportable_group: tuple[PaymentPlanGroup, Payment],
) -> None:
    group, _ = exportable_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-group"]')
        browser.click('[data-cy="button-delivery-export-xlsx-group"]')

        browser.wait_for_element_visible('[data-cy="dialog-delivery-export-xlsx-group"]')
        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-group-submit"]')
        browser.click('[data-cy="button-delivery-export-xlsx-group-submit"]')

        browser.wait_for_text("Export started")

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")
        browser.wait_for_text("Batch #1")
        browser.wait_for_element_visible('[data-cy="batch-download-link-1"]')


def test_export_payment_plan_group_with_auth_code(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    exportable_group: tuple[PaymentPlanGroup, Payment],
    auth_code_template: FinancialServiceProviderXlsxTemplate,
) -> None:
    group, _ = exportable_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
        Permissions.PM_EXPORT_XLSX_FOR_FSP,
        Permissions.PM_DOWNLOAD_FSP_AUTH_CODE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')
        browser.click('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')

        browser.wait_for_element_visible('[data-cy="dialog-delivery-export-xlsx-with-auth-code-group"]')
        # single exportable plan type -> shown as a locked field; the template picker
        # is the only enabled input in the dialog
        browser.wait_for_element_visible('[data-cy="locked-delivery-export-xlsx-with-auth-code-group-plan-type"]')
        template_input = browser.find_element(
            '[data-cy="dialog-delivery-export-xlsx-with-auth-code-group"] input:not([disabled])'
        )
        template_input.click()
        template_input.send_keys("Auth Code Template")
        browser.select_listbox_element("Auth Code Template")

        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-with-auth-code-group-submit"]')
        browser.click('[data-cy="button-delivery-export-xlsx-with-auth-code-group-submit"]')

        browser.wait_for_text("Export started")

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")
        browser.wait_for_text("Batch #1")
        browser.wait_for_element_visible('[data-cy="batch-download-link-1"]')


def test_import_payment_plan_group_reconciliation(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    exportable_group: tuple[PaymentPlanGroup, Payment],
    reconciliation_file: str,
) -> None:
    group, payment = exportable_group
    payment_plan = payment.parent
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_element_clickable('[data-cy="button-delivery-import-xlsx-group"]')
        browser.click('[data-cy="button-delivery-import-xlsx-group"]')

        browser.wait_for_element_visible('[data-cy="dialog-delivery-import-xlsx-group"]')
        browser.choose_file('[data-cy="dialog-delivery-import-xlsx-group"] input[type="file"]', reconciliation_file)

        browser.wait_for_element_clickable('[data-cy="button-delivery-import-xlsx-group-submit"]')
        browser.click('[data-cy="button-delivery-import-xlsx-group-submit"]')

        browser.wait_for_text("Delivery reconciliation import started")

        payment.refresh_from_db()
        payment_plan.refresh_from_db()
        assert payment.delivered_quantity == Decimal("50.00")
        assert payment.status == Payment.STATUS_DISTRIBUTION_PARTIAL
        assert payment_plan.status == PaymentPlan.Status.FINISHED


def _open_group_reconciliation_dialog(browser: HopeTestBrowser, file_path: str) -> None:
    browser.wait_for_element_clickable('[data-cy="button-delivery-import-xlsx-group"]')
    browser.click('[data-cy="button-delivery-import-xlsx-group"]')
    browser.wait_for_element_visible('[data-cy="dialog-delivery-import-xlsx-group"]')
    browser.choose_file('[data-cy="dialog-delivery-import-xlsx-group"] input[type="file"]', file_path)


def _enable_reconciliation_override(browser: HopeTestBrowser) -> None:
    dialog_selector = '[data-cy="dialog-delivery-import-xlsx-group"]'
    browser.wait_for_element_clickable(f"{dialog_selector} label")
    browser.click(f"{dialog_selector} label")
    browser.wait_for_element_visible(f'{dialog_selector} [role="combobox"]')


def _submit_group_reconciliation(browser: HopeTestBrowser) -> None:
    browser.wait_for_element_clickable('[data-cy="button-delivery-import-xlsx-group-submit"]')
    browser.click('[data-cy="button-delivery-import-xlsx-group-submit"]')
    browser.wait_for_text("Delivery reconciliation import started")


def test_override_reconciliation_resets_payment_for_empty_quantity(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    finished_group_with_empty_reconciliation_file: tuple[PaymentPlanGroup, Payment, str],
) -> None:
    group, payment, reconciliation_file_path = finished_group_with_empty_reconciliation_file
    payment_plan = payment.parent
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
        Permissions.PM_IMPORT_XLSX_WITH_RECONCILIATION_OVERRIDE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        _open_group_reconciliation_dialog(browser, reconciliation_file_path)
        _enable_reconciliation_override(browser)
        browser.wait_for_text(
            "Reset rows with empty/null delivered_quantity",
            '[data-cy="dialog-delivery-import-xlsx-group"] [role="combobox"]',
        )
        _submit_group_reconciliation(browser)

        payment.refresh_from_db()
        payment_plan.refresh_from_db()
        assert payment.delivered_quantity is None
        assert payment.delivered_quantity_usd is None
        assert payment.delivery_date is None
        assert payment.status == Payment.STATUS_SENT_TO_FSP
        assert payment.transaction_reference_id is None
        assert payment.reason_for_unsuccessful_payment is None
        assert payment.additional_collector_name is None
        assert payment.additional_document_type is None
        assert payment.additional_document_number is None
        assert payment.transaction_status_blockchain_link is None
        assert payment.extra_fields == {}
        assert payment.fsp_extra_fields == {"fsp_reference": "keep"}
        assert payment_plan.status == PaymentPlan.Status.ACCEPTED
        assert payment_plan.total_delivered_quantity == Decimal(0)


def test_override_reconciliation_ignores_payment_for_empty_quantity(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    finished_group_with_empty_reconciliation_file: tuple[PaymentPlanGroup, Payment, str],
) -> None:
    group, payment, reconciliation_file_path = finished_group_with_empty_reconciliation_file
    payment_plan = payment.parent
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
        Permissions.PM_IMPORT_XLSX_WITH_RECONCILIATION_OVERRIDE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        _open_group_reconciliation_dialog(browser, reconciliation_file_path)
        _enable_reconciliation_override(browser)
        browser.click('[data-cy="dialog-delivery-import-xlsx-group"] [role="combobox"]')
        browser.select_listbox_element("Ignore rows with empty/null delivered_quantity")
        _submit_group_reconciliation(browser)

        payment.refresh_from_db()
        payment_plan.refresh_from_db()
        assert payment.delivered_quantity == Decimal("50.00")
        assert payment.delivered_quantity_usd == Decimal("5.00")
        assert payment.status == Payment.STATUS_DISTRIBUTION_PARTIAL
        assert payment.transaction_reference_id == "ORIGINAL-REFERENCE"
        assert payment.reason_for_unsuccessful_payment == "Original reason"
        assert payment.additional_collector_name == "Original collector"
        assert payment.additional_document_type == "National ID"
        assert payment.additional_document_number == "ORIGINAL-DOCUMENT"
        assert payment.transaction_status_blockchain_link == "https://example.com/original-transaction"
        assert payment.extra_fields == {"reconciliation_note": "original"}
        assert payment.fsp_extra_fields == {"fsp_reference": "keep"}
        assert payment_plan.status == PaymentPlan.Status.FINISHED


def test_group_reconciliation_preserves_closed_plan_and_aborts_when_closed_quantity_changes(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    mixed_reconciliation_group: tuple[PaymentPlanGroup, PaymentPlan, PaymentPlan, Payment, Payment],
    mixed_reconciliation_file: str,
) -> None:
    group, first_plan, second_plan, first_payment, second_payment = mixed_reconciliation_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
        Permissions.PM_IMPORT_XLSX_WITH_RECONCILIATION_OVERRIDE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        _open_group_reconciliation_dialog(browser, mixed_reconciliation_file)
        _submit_group_reconciliation(browser)

        first_payment.refresh_from_db()
        second_payment.refresh_from_db()
        first_plan.refresh_from_db()
        second_plan.refresh_from_db()
        assert first_payment.delivered_quantity == Decimal("50.00")
        assert second_payment.delivered_quantity is None
        assert first_plan.status == PaymentPlan.Status.FINISHED
        assert second_plan.status == PaymentPlan.Status.ACCEPTED

        first_plan_flow = PaymentPlanFlow(first_plan)
        first_plan_flow.status_ready_for_closure()
        first_plan_flow.status_close()
        first_plan.save()
        _set_delivered_quantities(mixed_reconciliation_file, {str(second_payment.unicef_id): 75})
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        _open_group_reconciliation_dialog(browser, mixed_reconciliation_file)
        _submit_group_reconciliation(browser)

        first_payment.refresh_from_db()
        second_payment.refresh_from_db()
        first_plan.refresh_from_db()
        second_plan.refresh_from_db()
        assert first_payment.delivered_quantity == Decimal("50.00")
        assert second_payment.delivered_quantity == Decimal("75.00")
        assert first_plan.status == PaymentPlan.Status.CLOSED
        assert second_plan.status == PaymentPlan.Status.FINISHED

        _set_delivered_quantities(
            mixed_reconciliation_file,
            {
                str(first_payment.unicef_id): 60,
                str(second_payment.unicef_id): 80,
            },
        )
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        _open_group_reconciliation_dialog(browser, mixed_reconciliation_file)
        _enable_reconciliation_override(browser)
        browser.wait_for_element_clickable('[data-cy="button-delivery-import-xlsx-group-submit"]')
        browser.click('[data-cy="button-delivery-import-xlsx-group-submit"]')
        browser.wait_for_text("The entire file cannot be imported.")

        first_payment.refresh_from_db()
        second_payment.refresh_from_db()
        first_plan.refresh_from_db()
        second_plan.refresh_from_db()
        assert first_payment.delivered_quantity == Decimal("50.00")
        assert second_payment.delivered_quantity == Decimal("75.00")
        assert first_plan.status == PaymentPlan.Status.CLOSED
        assert second_plan.status == PaymentPlan.Status.FINISHED


def test_send_payment_plan_group_to_payment_gateway(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    sendable_group: tuple[PaymentPlanGroup, PaymentPlan],
) -> None:
    group, plan = sendable_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_SEND_TO_PAYMENT_GATEWAY,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_element_clickable('[data-cy="button-send-to-payment-gateway-group"]')

        with mock.patch("hope.apps.payment.services.payment_plan_services.send_to_payment_gateway_async_task"):
            browser.click('[data-cy="button-send-to-payment-gateway-group"]')
            browser.wait_for_text("Sending to Payment Gateway started")

        plan.refresh_from_db()
        assert plan.background_action_status == PaymentPlan.BackgroundActionStatus.SEND_TO_PAYMENT_GATEWAY


def test_group_shows_batch_with_download_link(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    group_with_exported_batch: PaymentPlanGroup,
) -> None:
    group = group_with_exported_batch
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_text("Batch #1")
        browser.wait_for_element_visible('[data-cy="batch-download-link-1"]')


def test_group_payment_plan_list_export_tag_links_to_batch(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    group_with_exported_batch: PaymentPlanGroup,
) -> None:
    group = group_with_exported_batch
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_text("Export Batch")
        browser.wait_for_element_clickable(f'table a[href$="/groups/{group.id}/batches/1"]')
        browser.click(f'table a[href$="/groups/{group.id}/batches/1"]')

        browser.wait_for_element_visible('[data-cy="button-download-batch"]')


def test_batch_detail_shows_download_button_when_file_present(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    group_with_exported_batch: PaymentPlanGroup,
) -> None:
    group = group_with_exported_batch
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}/batches/1")

        browser.wait_for_element_visible('[data-cy="button-download-batch"]')
        browser.assert_element_absent('[data-cy="button-export-batch"]')


def test_batch_detail_shows_reexport_button_when_file_missing(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    group_with_unexported_batch: PaymentPlanGroup,
) -> None:
    group = group_with_unexported_batch
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}/batches/1")

        browser.wait_for_element_visible('[data-cy="button-export-batch"]')
        browser.assert_element_absent('[data-cy="button-download-batch"]')


def test_linked_payment_plans_modal_on_cycle_details(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    plan_with_follow_up: tuple[PaymentPlan, PaymentPlan],
) -> None:
    source, follow_up = plan_with_follow_up
    cycle = source.program_cycle
    program = cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_VIEW_DETAILS,
        Permissions.PM_PROGRAMME_CYCLE_VIEW_DETAILS,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/program-cycles/{cycle.id}")

        _assert_linked_plans_modal_opens_and_closes(browser, source, follow_up)


def test_linked_payment_plans_modal_on_payment_plans_list(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    plan_with_follow_up: tuple[PaymentPlan, PaymentPlan],
) -> None:
    source, follow_up = plan_with_follow_up
    program = source.program_cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_VIEW_DETAILS,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/payment-plans")

        _assert_linked_plans_modal_opens_and_closes(browser, source, follow_up)


def test_group_details_action_buttons_follow_group_state(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    exportable_group: tuple[PaymentPlanGroup, Payment],
    payment_plan_group: PaymentPlanGroup,
) -> None:
    """With every permission granted, the visible actions depend only on group state.

    The two states are mutually exclusive: export needs an exportable plan, delete
    needs a group with none, so no single group can ever show all six buttons.
    """
    group, _ = exportable_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_UPDATE,
        Permissions.PM_PAYMENT_PLAN_GROUP_DELETE,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
        Permissions.PM_PAYMENT_PLAN_GROUP_SEND_TO_PAYMENT_GATEWAY,
        Permissions.PM_DOWNLOAD_FSP_AUTH_CODE,
    ):
        browser.login(username="noperm_user", password="testtest2")

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")
        browser.wait_for_element_visible('[data-cy="button-edit-group-name"]')
        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-group"]')
        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')
        browser.wait_for_element_visible('[data-cy="button-delivery-import-xlsx-group"]')
        browser.wait_for_element_visible('[data-cy="button-send-to-payment-gateway-group"]')
        browser.assert_element_absent('[data-cy="button-delete-group"]')

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{payment_plan_group.id}")
        browser.wait_for_element_visible('[data-cy="button-delete-group"]')
        browser.wait_for_element_visible('[data-cy="button-edit-group-name"]')
        browser.wait_for_element_visible('[data-cy="button-delivery-import-xlsx-group"]')
        browser.wait_for_element_visible('[data-cy="button-send-to-payment-gateway-group"]')
        browser.assert_element_absent('[data-cy="button-delivery-export-xlsx-group"]')
        browser.assert_element_absent('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')


def test_group_details_actions_are_pinned_to_their_permissions(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    exportable_group: tuple[PaymentPlanGroup, Payment],
) -> None:
    group, _ = exportable_group
    group_url = f"/{business_area.slug}/programs/{group.cycle.program.code}/payment-module/groups/{group.id}"

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
        Permissions.PM_PAYMENT_PLAN_GROUP_SEND_TO_PAYMENT_GATEWAY,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(group_url)

        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-group"]')
        browser.wait_for_element_visible('[data-cy="button-send-to-payment-gateway-group"]')
        browser.assert_element_absent('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')
        browser.assert_element_absent('[data-cy="button-edit-group-name"]')
        browser.assert_element_absent('[data-cy="button-delivery-import-xlsx-group"]')

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_UPDATE,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
    ):
        browser.open(group_url)

        browser.wait_for_element_visible('[data-cy="button-edit-group-name"]')
        browser.wait_for_element_visible('[data-cy="button-delivery-import-xlsx-group"]')
        browser.assert_element_absent('[data-cy="button-delivery-export-xlsx-group"]')
        browser.assert_element_absent('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]')
        browser.assert_element_absent('[data-cy="button-send-to-payment-gateway-group"]')


def test_group_details_xlsx_buttons_disabled_while_background_action_busy(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    busy_group: PaymentPlanGroup,
) -> None:
    group_url = f"/{business_area.slug}/programs/{busy_group.cycle.program.code}/payment-module/groups/{busy_group.id}"

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_UPDATE,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
        Permissions.PM_PAYMENT_PLAN_GROUP_IMPORT_XLSX,
        Permissions.PM_DOWNLOAD_FSP_AUTH_CODE,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(group_url)

        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-group"]:disabled')
        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]:disabled')
        browser.wait_for_element_visible('[data-cy="button-delivery-import-xlsx-group"]:disabled')
        # Renaming does not touch the XLSX pipeline, so it stays available.
        browser.assert_element_present('[data-cy="button-edit-group-name"]:not(:disabled)')

        # A running reconciliation import blocks the same actions as a running export.
        busy_group.background_action_status = PaymentPlanGroup.BackgroundActionStatus.XLSX_IMPORTING_RECONCILIATION
        busy_group.save(update_fields=["background_action_status"])
        browser.open(group_url)

        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-group"]:disabled')
        browser.wait_for_element_visible('[data-cy="button-delivery-export-xlsx-with-auth-code-group"]:disabled')
        browser.wait_for_element_visible('[data-cy="button-delivery-import-xlsx-group"]:disabled')


def test_group_details_shows_background_action_status_while_busy(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    busy_group: PaymentPlanGroup,
) -> None:
    group_url = f"/{business_area.slug}/programs/{busy_group.cycle.program.code}/payment-module/groups/{busy_group.id}"

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(group_url)

        browser.wait_for_element_visible('[data-cy="group-background-action-status"]')
        browser.assert_text("XLSX EXPORTING", '[data-cy="group-background-action-status"]')


def test_group_details_overview_shows_totals_and_links_to_cycle(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    group_with_totals: PaymentPlanGroup,
) -> None:
    cycle = group_with_totals.cycle
    program = cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PROGRAMME_CYCLE_VIEW_DETAILS,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group_with_totals.id}")

        # The serializer sums the plan totals and DRF encodes the Decimal as a float.
        browser.wait_for_text(group_with_totals.name, 'div[data-cy="label-Name"]')
        browser.wait_for_text("125.5", 'div[data-cy="label-Total Entitled (USD)"]')
        browser.wait_for_text("25.5", 'div[data-cy="label-Total Delivered (USD)"]')
        browser.wait_for_text("100", 'div[data-cy="label-Total Undelivered (USD)"]')

        browser.wait_for_element_clickable('div[data-cy="label-Cycle"] a').click()
        browser.wait_for_text(cycle.title, 'h5[data-cy="page-header-title"]')
        browser.wait_for_element_visible('div[data-cy="label-Frequency of Payment"]')


def test_group_details_batches_section_reflects_export_state(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    payment_plan_group: PaymentPlanGroup,
    group_with_unexported_batch: PaymentPlanGroup,
) -> None:
    """No batches at all hides the section; a tagged batch without a file lists it
    without a Download link.
    """
    program = payment_plan_group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
    ):
        browser.login(username="noperm_user", password="testtest2")

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{payment_plan_group.id}")
        browser.wait_for_text(payment_plan_group.name, 'h5[data-cy="page-header-title"]')
        browser.assert_element_absent('[data-cy="batches-section"]')

        browser.open(
            f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group_with_unexported_batch.id}"
        )
        browser.wait_for_text("Batch #1")
        browser.assert_element_absent('[data-cy="batch-download-link-1"]')


def test_group_details_export_dialog_selects_plan_type_when_group_has_several(
    browser: HopeTestBrowser,
    user_with_no_permissions: User,
    business_area: BusinessArea,
    multi_plan_type_group: tuple[PaymentPlanGroup, PaymentPlan, PaymentPlan],
) -> None:
    group, regular_plan, top_up_plan = multi_plan_type_group
    program = group.cycle.program

    with grant_permission(
        user_with_no_permissions,
        business_area,
        Permissions.PROGRAMME_VIEW_LIST_AND_DETAILS,
        Permissions.PM_VIEW_LIST,
        Permissions.PM_PAYMENT_PLAN_GROUP_VIEW_DETAIL,
        Permissions.PM_PAYMENT_PLAN_GROUP_EXPORT_XLSX,
    ):
        browser.login(username="noperm_user", password="testtest2")
        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")

        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-group"]').click()

        browser.wait_for_element_visible('[data-cy="dialog-delivery-export-xlsx-group"]')
        # Two exportable plan types turn the read-only locked field into a picker.
        browser.wait_for_element_visible('[data-cy="select-delivery-export-xlsx-group-plan-type"]')
        browser.assert_element_absent('[data-cy="locked-delivery-export-xlsx-group-plan-type"]')

        browser.click('[data-cy="select-delivery-export-xlsx-group-plan-type"] input')
        browser.select_listbox_element("Top Up")

        browser.wait_for_element_clickable('[data-cy="button-delivery-export-xlsx-group-submit"]').click()
        browser.wait_for_text("Export started")

        browser.open(f"/{business_area.slug}/programs/{program.code}/payment-module/groups/{group.id}")
        browser.wait_for_text("Batch #1 Top Up")

        # Only the picked plan type is batched; the regular plan stays unexported.
        top_up_plan.refresh_from_db()
        regular_plan.refresh_from_db()
        assert top_up_plan.export_tag == 1
        assert regular_plan.export_tag is None
