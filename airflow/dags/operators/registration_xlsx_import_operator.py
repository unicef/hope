from io import BytesIO

from django.contrib.gis.geos import Point
from django.core.files import File
from django.core.files.storage import default_storage
from django.db import transaction
from django.utils import timezone
from openpyxl_image_loader import SheetImageLoader

from .base import DjangoOperator


class RegistrationXLSXImportOperator(DjangoOperator):
    """
    Works on valid XLSX files, parsing them and creating households/individuals
    in the Registration Datahub. Once finished it will update the status of
    that registration data import instance.
    """

    image_loader = None
    business_area = None
    households = None
    individuals = None
    documents = None

    def _handle_document_fields(
        self, value, header, individual, *args, **kwargs
    ):
        if value is None:
            return

        if header == "other_id_type_i_c":
            doc_type = value
        else:
            readable_name = (
                header.replace("_no", "").replace("_i_c", "").replace("_", " ")
            )
            readable_name_split = readable_name.split(" ")
            doc_type = ""
            for word in readable_name_split:
                if word == "id":
                    doc_type += word.upper() + " "
                else:
                    doc_type += word.capitalize() + " "
            doc_type = doc_type.strip()
        row_num = kwargs.get("row_number")
        document_data = self.documents.get(f"individual_{row_num}")
        if document_data:
            document_data["value"] = value
        else:
            self.documents[f"individual_{row_num}"] = {
                "individual": individual,
                "header": header,
                "type": doc_type,
                "value": value,
            }

    def _handle_document_photo_fields(
        self, cell, header, row_num, individual, *args, **kwargs
    ):
        if not self.image_loader.image_in(cell.coordinate):
            return

        document_data = self.documents.get(f"individual_{row_num}")

        image = self.image_loader.get(cell.coordinate)
        file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
        file_io = BytesIO()

        image.save(file_io, image.format)

        file = File(file_io, name=file_name)

        if document_data:
            document_data["photo"] = file
        else:
            self.documents[f"individual_{row_num}"] = {
                "individual": individual,
                "header": header,
                "photo": file_name,
            }

    def _handle_image_field(
        self, cell, is_flex_field=False, *args, **kwargs
    ):
        if self.image_loader.image_in(cell.coordinate):
            image = self.image_loader.get(cell.coordinate)
            file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
            file_io = BytesIO()

            image.save(file_io, image.format)

            file = File(file_io, name=file_name)

            if is_flex_field:
                return default_storage.save(file_name, file)

            return file

    def _handle_geopoint_field(self, value, header, *args, **kwargs):
        if not value:
            return ""

        values_as_list = value.split(",")
        longitude = values_as_list[0].strip()
        latitude = values_as_list[1].strip()
        return Point(x=float(longitude), y=float(latitude), srid=4326)

    def _create_documents(self):
        from django_countries.fields import Country
        from household.const import COUNTRIES_NAME_ALPHA2
        from registration_datahub.models import (
            ImportedDocument,
            ImportedDocumentType,
        )

        docs_to_create = []
        for document_data in self.documents.values():
            doc_type = ImportedDocumentType.objects.filter(
                country=Country(
                    COUNTRIES_NAME_ALPHA2.get(
                        self.business_area.name.capitalize()
                    )
                ),
                label=document_data.get("type"),
            ).first()
            photo = document_data.get("photo")
            individual = document_data.get("individual")
            obj = ImportedDocument(
                document_number=document_data.get("value"),
                photo=photo,
                individual=individual,
                type=doc_type,
            )

            docs_to_create.append(obj)

        ImportedDocument.objects.bulk_create(docs_to_create)

    def _create_objects(self, sheet, registration_data_import):
        from core.utils import (
            get_combined_attributes,
            serialize_flex_attributes,
        )
        from registration_datahub.models import (
            ImportedHousehold,
            ImportedIndividual,
        )

        complex_fields = {
            "individuals": {
                "birth_certificate_no_i_c": self._handle_document_fields,
                "birth_certificate_photo_i_c": self._handle_document_photo_fields,
                "drivers_license_no_i_c": self._handle_document_fields,
                "drivers_license_photo_i_c": self._handle_document_photo_fields,
                "electoral_card_no_i_c": self._handle_document_fields,
                "electoral_card_photo_i_c": self._handle_document_photo_fields,
                "unhcr_id_no_i_c": self._handle_document_fields,
                "unhcr_id_photo_i_c": self._handle_document_photo_fields,
                "national_id_no_ic": self._handle_document_fields,
                "national_id_photo_ic": self._handle_document_photo_fields,
                "national_passport_i_c": self._handle_document_fields,
                "national_passport_photo_i_c": self._handle_document_photo_fields,
                "scope_id_no_i_c": self._handle_document_fields,
                "scope_id_photo_i_c": self._handle_document_photo_fields,
                "other_id_type_i_c": self._handle_document_fields,
                "other_id_no_i_c": self._handle_document_fields,
                "other_id_photo_i_c": self._handle_document_photo_fields,
                "photo_i_c": self._handle_image_field,
            },
            "households": {
                "consent_h_c": self._handle_image_field,
                "hh_geopoint_h_c": self._handle_geopoint_field,
            },
        }

        complex_types = {
            "GEOPOINT": self._handle_geopoint_field,
            "IMAGE": self._handle_image_field,
        }

        sheet_title = sheet.title.lower()

        combined_fields = get_combined_attributes()
        flex_fields = serialize_flex_attributes()

        first_row = sheet[1]
        households_to_update = []
        for row in sheet.iter_rows(min_row=3):
            if not any([cell.value for cell in row]):
                continue

            if sheet_title == "households":
                obj_to_create = ImportedHousehold(
                    registration_data_import=registration_data_import,
                )
            else:
                obj_to_create = ImportedIndividual(
                    registration_data_import=registration_data_import,
                )

            household_id = None
            for cell, header_cell in zip(row, first_row):
                header = header_cell.value
                current_field = combined_fields.get(header)

                if not current_field:
                    continue

                is_not_required_and_empty = (
                    not current_field.get("required") and cell.value is None
                )

                if is_not_required_and_empty:
                    continue

                if header == "household_id":
                    household_id = cell.value
                    if sheet_title == "individuals":
                        obj_to_create.household_id = self.households.get(
                            household_id
                        ).pk

                if header in complex_fields[sheet_title]:
                    fn = complex_fields[sheet_title].get(header)
                    value = fn(
                        value=cell.value,
                        cell=cell,
                        header=header,
                        row_num=cell.row,
                        individual=obj_to_create
                        if sheet_title == "individuals"
                        else None,
                    )
                    if value is not None:
                        setattr(
                            obj_to_create,
                            combined_fields[header]["name"],
                            value,
                        )
                elif (
                    hasattr(obj_to_create, combined_fields[header]["name"],)
                    and header != "household_id"
                ):
                    value = cell.value
                    if not value:
                        continue
                    if header == "relationship_i_c" and value == "HEAD":
                        household = self.households.get(household_id)
                        household.head_of_household = obj_to_create
                        households_to_update.append(household)

                    setattr(
                        obj_to_create,
                        combined_fields[header]["name"],
                        value,
                    )
                elif header in flex_fields[sheet_title]:
                    value = cell.value
                    type_name = flex_fields[sheet_title][header]["type"]
                    if type_name in complex_types:
                        fn = complex_types[type_name]
                        value = fn(
                            value=cell.value,
                            cell=cell,
                            header=header,
                            is_flex_field=True,
                        )
                    obj_to_create.flex_fields[header] = value

            if sheet_title == "households":
                self.households[household_id] = obj_to_create
            else:
                self.individuals.append(obj_to_create)

        if sheet_title == "households":
            ImportedHousehold.objects.bulk_create(self.households.values())
        else:
            ImportedIndividual.objects.bulk_create(self.individuals)
            ImportedHousehold.objects.bulk_update(
                households_to_update, ["head_of_household"], 1000,
            )
            self._create_documents()

    @transaction.atomic()
    def execute(self, context, **kwargs):
        import openpyxl

        from core.models import BusinessArea
        from registration_datahub.models import RegistrationDataImportDatahub
        from registration_datahub.models import ImportData

        self.households = {}
        self.documents = {}
        self.individuals = []

        dag_run = context["dag_run"]
        config_vars = dag_run.conf

        registration_data_import = RegistrationDataImportDatahub.objects.get(
            id=config_vars.get("registration_data_import_id")
        )
        import_data = ImportData.objects.get(
            id=config_vars.get("import_data_id"),
        )

        business_area_id = config_vars.get("business_area")
        self.business_area = BusinessArea.objects.get(id=business_area_id)

        wb = openpyxl.load_workbook(import_data.xlsx_file, data_only=True)

        # households objects have to be create first
        worksheets = (wb["Households"], wb["Individuals"])
        for sheet in worksheets:
            self.image_loader = SheetImageLoader(sheet)
            self._create_objects(sheet, registration_data_import)

        registration_data_import.import_done = True
        registration_data_import.save()
