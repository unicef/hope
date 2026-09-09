from __future__ import annotations

import datetime
from decimal import Decimal, InvalidOperation
import logging
from typing import IO, TYPE_CHECKING, Any, cast

from dateutil.parser import parse
from django.contrib.contenttypes.models import ContentType
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
import openpyxl
import pytz

from hope.apps.activity_log.utils import copy_model_object
from hope.apps.grievance.models import GrievanceTicket, TicketPaymentVerificationDetails
from hope.apps.payment.services.handle_total_cash_in_households import (
    handle_total_cash_in_specific_households,
)
from hope.apps.payment.utils import (
    bulk_log_payment_changes,
    calculate_counts,
    get_payment_delivered_quantity_status_and_value,
    get_quantity_in_usd,
)
from hope.apps.payment.xlsx.base_xlsx_import_service import XlsxImportBaseService
from hope.apps.payment.xlsx.xlsx_error import XlsxError
from hope.models import (
    FileTemp,
    FinancialServiceProviderXlsxTemplate,
    Payment,
    PaymentVerification,
    PaymentVerificationPlan,
    User,
)

if TYPE_CHECKING:
    import io

    from django.db.models import QuerySet
    from openpyxl.cell.cell import Cell

    from hope.models import PaymentPlan


class XlsxPaymentPlanDeliveryImportService(XlsxImportBaseService):
    logger = logging.getLogger(__name__)
    KNOWN_COLUMNS: frozenset[str] = frozenset(FinancialServiceProviderXlsxTemplate.DEFAULT_COLUMNS)
    NULL_DELIVERY_POLICIES = frozenset({"ignore", "reset"})

    ACTION_APPLY = "apply"
    ACTION_CONFLICT = "conflict"
    ACTION_IGNORE = "ignore"
    ACTION_RESET = "reset"
    ACTION_SKIP = "skip"

    OVERRIDE_ELIGIBLE_STATUSES = frozenset(
        {
            Payment.STATUS_SENT_TO_FSP,
            Payment.STATUS_DISTRIBUTION_SUCCESS,
            Payment.STATUS_DISTRIBUTION_PARTIAL,
            Payment.STATUS_NOT_DISTRIBUTED,
            Payment.STATUS_ERROR,
        }
    )
    PAYMENT_UPDATE_FIELDS = (
        "delivered_quantity",
        "delivered_quantity_usd",
        "status",
        "status_date",
        "delivery_date",
        "transaction_reference_id",
        "reason_for_unsuccessful_payment",
        "additional_collector_name",
        "additional_document_type",
        "additional_document_number",
        "transaction_status_blockchain_link",
        "extras",
    )
    OPTIONAL_RECONCILIATION_FIELDS = {
        "reference_id": "transaction_reference_id",
        "reason_for_unsuccessful_payment": "reason_for_unsuccessful_payment",
        "additional_collector_name": "additional_collector_name",
        "additional_document_type": "additional_document_type",
        "additional_document_number": "additional_document_number",
        "transaction_status_blockchain_link": "transaction_status_blockchain_link",
    }
    MISSING = object()

    class XlsxPaymentPlanDeliveryImportServiceError(Exception):
        pass

    def __init__(
        self,
        payment_plan: "PaymentPlan",
        file: io.BytesIO | IO[bytes],
        fsp_owned_headers: set[str] | None = None,
        override: bool = False,
        null_delivery_policy: str = "reset",
    ) -> None:
        if null_delivery_policy not in self.NULL_DELIVERY_POLICIES:
            raise ValueError(f"Unsupported null delivery policy: {null_delivery_policy}")

        self.payment_plan = payment_plan
        self.pp_currency_exchange_date = self.payment_plan.currency_exchange_date
        self.payment_list: QuerySet[Payment] = payment_plan.eligible_payments.select_related("household")
        self.file = file
        self.errors: list[XlsxError] = []
        self.conflict_errors: list[XlsxError] = []
        self.payments_dict: dict[str, Payment] = {str(payment.unicef_id): payment for payment in self.payment_list}
        self.fsp_owned_headers = (
            fsp_owned_headers
            if fsp_owned_headers is not None
            else {header for payment in self.payments_dict.values() for header in payment.fsp_extra_fields}
        )
        self.override = override
        self.null_delivery_policy = null_delivery_policy
        self.payment_ids: list[str] = list(self.payments_dict)
        self.payment_ids_from_xlsx: list[str] = []
        self.payments_to_save: list[Payment] = []
        self.old_payments: dict[Any, Payment] = {}
        self.payment_ids_for_verification_cleanup: set[Any] = set()
        self.skipped_rows: list[dict[str, Any]] = []
        self.required_columns: list[str] = ["payment_id", "delivered_quantity"]
        self.xlsx_headers: list[str] = []
        self.is_updated = False

    def open_workbook(self) -> openpyxl.Workbook:
        self.logger.info(f"Opening workbook for payment plan: {self.payment_plan.id}")
        wb = openpyxl.load_workbook(cast("Any", self.file), data_only=True)
        self.wb = wb
        self.ws_payments = wb[wb.sheetnames[0]]
        self.sheetname = wb.sheetnames[0]
        self.logger.info("Generating headers")
        self.xlsx_headers = [header.value for header in self.ws_payments[1]]
        return wb

    def _validate_headers(self) -> None:
        for required_column in self.required_columns:
            if required_column not in self.xlsx_headers:
                self.errors.append(
                    XlsxError(
                        self.sheetname,
                        None,
                        f"Provided headers {self.xlsx_headers} do not match expected headers. "
                        f"{self.required_columns} are required headers.",
                    )
                )
                return

    def _validate_payment_id(self, row: tuple[Cell, ...]) -> None:
        cell = row[self.xlsx_headers.index("payment_id")]
        payment_id = str(cell.value) if cell.value is not None else None
        if payment_id not in self.payment_ids:
            self.errors.append(
                XlsxError(
                    self.sheetname,
                    cell.coordinate,
                    f"This payment id {cell.value} is not in Payment Plan Payment List",
                )
            )
        if payment_id in self.payment_ids_from_xlsx:
            self.errors.append(
                XlsxError(
                    self.sheetname,
                    cell.coordinate,
                    f"Payment id {cell.value} appears multiple times in the import file",
                )
            )
        if payment_id is not None:
            self.payment_ids_from_xlsx.append(payment_id)

    def _payment_plan_is_eligible(self) -> bool:
        if self.payment_plan.status == self.payment_plan.Status.CLOSED:
            return False
        if self.override:
            return self.payment_plan.status in {
                self.payment_plan.Status.ACCEPTED,
                self.payment_plan.Status.FINISHED,
            }
        return self.payment_plan.status == self.payment_plan.Status.ACCEPTED

    def _should_skip_row(self, payment_id: str | None) -> bool:
        if payment_id is None or not self._payment_plan_is_eligible():
            return True
        payment = self.payments_dict.get(str(payment_id))
        if payment is None:
            return True
        if self.override:
            return payment.status not in self.OVERRIDE_ELIGIBLE_STATUSES
        return payment.delivered_quantity is None and payment.status != Payment.STATUS_SENT_TO_FSP

    @staticmethod
    def _is_empty_quantity(value: Any) -> bool:
        return value is None or isinstance(value, str) and not value.strip()

    def _parse_delivered_quantity(self, value: Any) -> Decimal | None:
        if self._is_empty_quantity(value):
            return None
        if isinstance(value, bool | datetime.date):
            raise ValueError
        try:
            quantity = Decimal(str(value).strip())
        except (InvalidOperation, TypeError, ValueError):
            raise ValueError from None
        if not quantity.is_finite() or quantity < 0 and quantity != Decimal(-1):
            raise ValueError
        try:
            return quantity.quantize(Decimal("0.01"))
        except InvalidOperation:
            raise ValueError from None

    def _get_row_action(self, payment: Payment, delivered_quantity: Decimal | None) -> str:
        action = self.ACTION_SKIP
        if not self._payment_plan_is_eligible():
            pass
        elif self.override:
            if payment.status not in self.OVERRIDE_ELIGIBLE_STATUSES:
                pass
            elif delivered_quantity is None:
                action = self.ACTION_RESET if self.null_delivery_policy == "reset" else self.ACTION_IGNORE
            else:
                action = self.ACTION_APPLY
        elif delivered_quantity is None:
            action = self.ACTION_IGNORE
        elif payment.delivered_quantity is not None:
            if delivered_quantity == payment.delivered_quantity:
                action = self.ACTION_IGNORE
            else:
                action = self.ACTION_CONFLICT
        elif payment.status == Payment.STATUS_SENT_TO_FSP:
            action = self.ACTION_APPLY
        return action

    def _validate_delivered_quantity(self, row: tuple[Cell, ...]) -> tuple[str, Decimal | None]:
        payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
        payment = self.payments_dict.get(payment_id)
        if payment is None or not self._payment_plan_is_eligible():
            return self.ACTION_SKIP, None
        if self.override and payment.status not in self.OVERRIDE_ELIGIBLE_STATUSES:
            return self.ACTION_SKIP, None
        if not self.override and payment.delivered_quantity is None and payment.status != Payment.STATUS_SENT_TO_FSP:
            return self.ACTION_SKIP, None

        cell = row[self.xlsx_headers.index("delivered_quantity")]
        try:
            delivered_quantity = self._parse_delivered_quantity(cell.value)
        except ValueError:
            self.errors.append(
                XlsxError(
                    self.sheetname,
                    cell.coordinate,
                    f"Payment {payment_id}: Delivered quantity {cell.value} must be a number greater than or equal "
                    "to zero, exactly -1, or empty.",
                )
            )
            return self.ACTION_SKIP, None

        action = self._get_row_action(payment, delivered_quantity)
        if action == self.ACTION_CONFLICT:
            conflict_error = XlsxError(
                self.sheetname,
                cell.coordinate,
                f"Payment {payment_id}: Delivered quantity {delivered_quantity} conflicts with the existing "
                f"delivered quantity {payment.delivered_quantity}.",
            )
            self.errors.append(conflict_error)
            self.conflict_errors.append(conflict_error)
            return action, delivered_quantity

        if delivered_quantity is not None and delivered_quantity != Decimal("-1.00"):
            entitlement_quantity = payment.entitlement_quantity or Decimal(0)
            if delivered_quantity > entitlement_quantity:
                self.errors.append(
                    XlsxError(
                        self.sheetname,
                        cell.coordinate,
                        f"Payment {payment_id}: Delivered quantity {delivered_quantity} is bigger than "
                        f"Entitlement quantity {entitlement_quantity}",
                    )
                )
        return action, delivered_quantity

    def _validate_delivery_date(self, row: tuple[Cell, ...]) -> None:
        if "delivery_date" not in self.xlsx_headers:
            return
        payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
        if payment_id not in self.payments_dict:
            return
        cell = row[self.xlsx_headers.index("delivery_date")]
        if cell.value in (None, ""):
            return
        try:
            delivery_date = self._parse_delivery_date(cell.value)
            date_value = delivery_date.date()
            if date_value > datetime.date.today() or date_value < self.payment_plan.program.start_date:
                self.errors.append(
                    XlsxError(
                        self.sheetname,
                        cell.coordinate,
                        f"Payment {payment_id}: Delivery date ({date_value}) cannot be greater than today's date,"
                        " and cannot be before Programme's start date",
                    )
                )
        except (ValueError, TypeError, OverflowError) as error:
            self.errors.append(
                XlsxError(
                    self.sheetname,
                    cell.coordinate,
                    f"Payment {payment_id}: Delivered date {cell.value} is not a datetime. {error}",
                )
            )

    def _validate_reason_for_unsuccessful_payment(self, row: tuple[Cell, ...]) -> None:
        self._mark_updated_when_optional_value_differs(row, "reason_for_unsuccessful_payment")

    def _validate_reference_id(self, row: tuple[Cell, ...]) -> None:
        self._mark_updated_when_optional_value_differs(row, "reference_id", "transaction_reference_id")

    def _validate_extras(self, row: tuple[Cell, ...]) -> None:
        payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
        payment = self.payments_dict.get(payment_id)
        if payment and self._get_extras_for_row(row, payment.extra_fields) != payment.extra_fields:
            self.is_updated = True

    def _mark_updated_when_optional_value_differs(
        self,
        row: tuple[Cell, ...],
        header_name: str,
        field_name: str | None = None,
    ) -> None:
        payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
        payment = self.payments_dict.get(payment_id)
        if payment is None or header_name not in self.xlsx_headers:
            return
        value = self._get_optional_cell_value(row, header_name)
        if value != getattr(payment, field_name or header_name):
            self.is_updated = True

    def _validate_rows(self) -> None:
        self.is_updated = False
        self.payment_ids_from_xlsx = []
        self.skipped_rows = []
        self.conflict_errors = []
        for row in self.ws_payments.iter_rows(min_row=2):
            if not any(cell.value for cell in row):
                continue

            error_count = len(self.errors)
            self._validate_payment_id(row)
            if len(self.errors) != error_count:
                continue
            action, _ = self._validate_delivered_quantity(row)
            if action == self.ACTION_APPLY:
                self._validate_delivery_date(row)
            if action in {self.ACTION_APPLY, self.ACTION_RESET}:
                self.is_updated = True
            elif action == self.ACTION_SKIP:
                payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
                payment = self.payments_dict[payment_id]
                self.skipped_rows.append(
                    {
                        "row": getattr(row[0], "row", None),
                        "payment_id": payment_id,
                        "reason": self._get_skip_reason(payment),
                    }
                )

    def _get_skip_reason(self, payment: Payment) -> str:
        if not self._payment_plan_is_eligible():
            return f"Payment Plan status {self.payment_plan.status} is not eligible for this import mode."
        return f"Payment status {payment.status} is not eligible for this import mode."

    def _validate_imported_file(self) -> None:
        if not self.is_updated:
            self.errors.append(
                XlsxError(
                    self.sheetname,
                    None,
                    "There aren't any updates in imported file, please add changes and try again",
                )
            )

    def validate(self) -> None:
        self.logger.info("Starting validation")
        self._validate_headers()
        if not self.errors:
            self._validate_rows()
            self._validate_imported_file()
        self.logger.info("Finished validation")

    def import_payment_list(self, user_id: str | None = None) -> None:
        self.logger.info("Starting importing payment list")
        exchange_rate = self.payment_plan.exchange_rate
        self.payments_to_save = []
        self.old_payments = {}
        self.payment_ids_for_verification_cleanup = set()

        for row in self.ws_payments.iter_rows(min_row=2):
            if any(cell.value for cell in row):
                self._import_row(row, exchange_rate)

        Payment.objects.bulk_update(self.payments_to_save, self.PAYMENT_UPDATE_FIELDS, batch_size=500)
        user = User.objects.filter(pk=user_id).first() if user_id else None
        bulk_log_payment_changes(
            [(self.old_payments[payment.pk], payment) for payment in self.payments_to_save],
            user,
        )
        handle_total_cash_in_specific_households([payment.household_id for payment in self.payments_to_save])

        self._cleanup_payment_verifications()
        self.logger.info("Finished import payment list")

    def _get_delivered_quantity_status_and_value(
        self,
        delivered_quantity: int | float | str | Decimal,
        entitlement_quantity: Decimal,
        payment_id: str,
    ) -> tuple[str, Decimal | None]:
        try:
            parsed_quantity = self._parse_delivered_quantity(delivered_quantity)
            if parsed_quantity is None:
                raise ValueError
            return get_payment_delivered_quantity_status_and_value(str(parsed_quantity), entitlement_quantity)
        except ValueError:
            raise self.XlsxPaymentPlanDeliveryImportServiceError(
                f"Invalid delivered_quantity {delivered_quantity} provided for payment_id {payment_id}"
            ) from None

    def _get_optional_cell_value(self, row: tuple[Cell, ...], header_name: str) -> Any:
        if header_name not in self.xlsx_headers:
            return None
        value = row[self.xlsx_headers.index(header_name)].value
        return None if value == "" else value

    def _get_optional_cell_value_or_missing(self, row: tuple[Cell, ...], header_name: str) -> Any:
        if header_name not in self.xlsx_headers:
            return self.MISSING
        return self._get_optional_cell_value(row, header_name)

    def _cleanup_payment_verifications(self) -> None:
        if not self.payment_ids_for_verification_cleanup:
            return
        verifications = PaymentVerification.objects.filter(payment_id__in=self.payment_ids_for_verification_cleanup)
        plans = list(
            PaymentVerificationPlan.objects.select_for_update()
            .filter(pk__in=verifications.values("payment_verification_plan_id"))
            .order_by("pk")
        )
        ticket_ids = TicketPaymentVerificationDetails.objects.filter(
            Q(payment_verification__in=verifications) | Q(payment_verifications__in=verifications)
        ).values("ticket_id")
        # Delete tickets first: deleting a verification alone would only null its ticket's FK.
        GrievanceTicket.objects.filter(pk__in=ticket_ids).delete()
        verifications.delete()
        for plan in plans:
            remaining_count = plan.payment_record_verifications.count()
            if remaining_count:
                calculate_counts(plan)
                plan.sample_size = remaining_count
                plan.save(
                    update_fields=[
                        "sample_size",
                        "responded_count",
                        "received_count",
                        "not_received_count",
                        "received_with_problems_count",
                    ]
                )
            else:
                files = FileTemp.objects.filter(
                    content_type=ContentType.objects.get_for_model(PaymentVerificationPlan), object_id=str(plan.pk)
                )
                for file_temp in files:
                    if file_temp.file:
                        transaction.on_commit(lambda file=file_temp.file: file.delete(save=False), robust=True)
                files.delete()
                plan.delete()

    def _parse_delivery_date(self, value: Any) -> datetime.datetime:
        if isinstance(value, str):
            value = parse(value)
        elif isinstance(value, datetime.date) and not isinstance(value, datetime.datetime):
            value = datetime.datetime.combine(value, datetime.time.min)
        if not isinstance(value, datetime.datetime):
            raise TypeError("value is not a date")
        if value.tzinfo is None:
            value = pytz.utc.localize(value)
        return value

    def _set_payment_delivery_date(self, delivery_date: Any, payment: Payment) -> tuple[Any, Any]:
        parsed_delivery_date = self._parse_delivery_date(delivery_date) if delivery_date not in (None, "") else None
        payment_delivery_date = payment.delivery_date
        if payment_delivery_date:
            payment_delivery_date = payment_delivery_date.replace(tzinfo=None)
        return parsed_delivery_date, payment_delivery_date

    def _normalize_delivery_date(self, delivery_date: Any, payment_delivery_date: Any) -> Any:
        delivery_date = delivery_date.date() if isinstance(delivery_date, datetime.datetime) else delivery_date
        if delivery_date and (
            delivery_date > datetime.date.today() or delivery_date < self.payment_plan.program.start_date
        ):
            return payment_delivery_date
        return delivery_date

    def _payment_changed(self, old_payment: Payment, payment: Payment) -> bool:
        return any(getattr(old_payment, field) != getattr(payment, field) for field in self.PAYMENT_UPDATE_FIELDS)

    def _import_row(self, row: tuple[Cell, ...], exchange_rate: Decimal | float | None) -> None:
        payment_id = str(row[self.xlsx_headers.index("payment_id")].value)
        if self._should_skip_row(payment_id):
            return
        payment = self.payments_dict.get(payment_id)
        if payment is None:
            return
        try:
            delivered_quantity = self._parse_delivered_quantity(
                row[self.xlsx_headers.index("delivered_quantity")].value
            )
        except ValueError as error:
            raise self.XlsxPaymentPlanDeliveryImportServiceError(
                f"Invalid delivered_quantity provided for payment_id {payment_id}"
            ) from error
        action = self._get_row_action(payment, delivered_quantity)
        if action == self.ACTION_CONFLICT:
            raise self.XlsxPaymentPlanDeliveryImportServiceError(
                f"Delivered quantity conflict for payment_id {payment_id}"
            )
        if action not in {self.ACTION_APPLY, self.ACTION_RESET}:
            return

        old_payment = cast("Payment", copy_model_object(payment))
        if action == self.ACTION_RESET:
            self._reset_payment(payment)
            self.payment_ids_for_verification_cleanup.add(payment.pk)
        else:
            self._apply_reconciliation_values(payment, row, delivered_quantity, exchange_rate)
            if self.override and old_payment.delivered_quantity != payment.delivered_quantity:
                self.payment_ids_for_verification_cleanup.add(payment.pk)

        if not self._payment_changed(old_payment, payment):
            return
        self.old_payments[payment.pk] = old_payment
        self.payments_to_save.append(payment)

    def _reset_payment(self, payment: Payment) -> None:
        old_status = payment.status
        payment.delivered_quantity = None
        payment.delivered_quantity_usd = None
        payment.delivery_date = None
        payment.status = Payment.STATUS_SENT_TO_FSP
        if old_status != payment.status:
            payment.status_date = timezone.now()
        payment.transaction_reference_id = None
        payment.reason_for_unsuccessful_payment = None
        payment.additional_collector_name = None
        payment.additional_document_type = None
        payment.additional_document_number = None
        payment.transaction_status_blockchain_link = None
        payment.set_extra_fields({})

    def _apply_reconciliation_values(
        self,
        payment: Payment,
        row: tuple[Cell, ...],
        delivered_quantity: Decimal | None,
        exchange_rate: Decimal | float | None,
    ) -> None:
        if delivered_quantity is None:
            raise self.XlsxPaymentPlanDeliveryImportServiceError("Cannot apply an empty delivered quantity")
        status, stored_quantity = self._get_delivered_quantity_status_and_value(
            delivered_quantity,
            payment.entitlement_quantity or Decimal(0),
            str(payment.unicef_id),
        )
        old_status = payment.status
        payment.delivered_quantity = stored_quantity
        payment.delivered_quantity_usd = get_quantity_in_usd(
            amount=stored_quantity,
            currency=self.payment_plan.currency,
            exchange_rate=Decimal(exchange_rate or 0),
            currency_exchange_date=self.pp_currency_exchange_date,
        )
        payment.status = status
        if old_status != status:
            payment.status_date = timezone.now()

        delivery_date = self._get_optional_cell_value_or_missing(row, "delivery_date")
        if delivery_date is not self.MISSING:
            payment.delivery_date = self._parse_delivery_date(delivery_date) if delivery_date is not None else None

        for header, field_name in self.OPTIONAL_RECONCILIATION_FIELDS.items():
            value = self._get_optional_cell_value_or_missing(row, header)
            if value is not self.MISSING:
                setattr(payment, field_name, value)
        payment.set_extra_fields(self._get_extras_for_row(row, payment.extra_fields))

    def _get_values_for_update(self, row: tuple[Cell, ...]) -> tuple[Any, Any, Any, Any, Any, Any, Any]:
        return (
            self._get_optional_cell_value(row, "additional_collector_name"),
            self._get_optional_cell_value(row, "additional_document_number"),
            self._get_optional_cell_value(row, "additional_document_type"),
            self._get_optional_cell_value(row, "delivery_date"),
            self._get_optional_cell_value(row, "reason_for_unsuccessful_payment"),
            self._get_optional_cell_value(row, "reference_id"),
            self._get_optional_cell_value(row, "transaction_status_blockchain_link"),
        )

    def _get_additional_doc_values(self, row: tuple[Cell, ...]) -> tuple[Any, Any]:
        return (
            self._get_optional_cell_value(row, "additional_document_number"),
            self._get_optional_cell_value(row, "additional_document_type"),
        )

    def _get_extras_for_row(self, row: tuple[Cell, ...], current_extras: dict[str, object] | None = None) -> dict:
        extras = dict(current_extras or {})
        for idx, header in enumerate(self.xlsx_headers):
            if not header or header in self.KNOWN_COLUMNS or header in self.fsp_owned_headers:
                continue
            value = row[idx].value
            if value is None or value == "":
                extras.pop(header, None)
                continue
            if isinstance(value, Decimal):
                value = float(value)
            elif isinstance(value, datetime.datetime):
                value = value.isoformat()
            elif not isinstance(value, int | float | bool):
                value = str(value)
            extras[header] = value
        return extras
